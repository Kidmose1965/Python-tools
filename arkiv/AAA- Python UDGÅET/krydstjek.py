#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Krydstjek - validering af krydshenvisninger i kontrakter og bilag
==================================================================
1. Kortlægger alle definitioner (nummererede afsnit/punkter + bilag/appendiks)
2. Finder alle krydshenvisninger (jf. afsnit 4.2, se Bilag 3, iht. punkt 2.1 ...)
3. Validerer hver henvisning mod dokumentsamlingen

Resultat: Excel-rapport med ❌ UGYLDIGE, ⚠️ USIKRE og ✅ GYLDIGE henvisninger
samt en opsummering.

Kan køres alene:
    python krydstjek.py kontrakt.docx "Bilag 1 Tidsplan.docx" [...] -o rapport.xlsx
eller via knappen "Validér krydshenvisninger" i extractor_gui.py.

Kræver: extractor.py i samme mappe + openpyxl.
"""

import argparse
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

import extractor as ex

THIN = Border(*[Side(style="thin")] * 4)
FILLS = {"ugyldig": PatternFill("solid", start_color="F8CBAD"),
         "usikker": PatternFill("solid", start_color="FFE699"),
         "gyldig": PatternFill("solid", start_color="C6E0B4")}
SYMBOL = {"ugyldig": "❌ UGYLDIG", "usikker": "⚠️ USIKKER", "gyldig": "✅ GYLDIG"}

# -- mønstre ----------------------------------------------------------------
SEKTIONSORD = r"(?:afsnit|punkt|pkt\.?|kapitel|sektion|klausul|§)"
BILAGSORD = r"(?:underbilag|bilag|appendiks|appendix)"
NUM = r"\d+(?:\.\d+)*"

RE_BILAG_SEKTION = re.compile(
    rf"\b({BILAGSORD})\s+(\d+[a-zA-Z]?|[A-ZÆØÅ])\b[,\s]+\s*({SEKTIONSORD})\s*({NUM})",
    re.IGNORECASE)
RE_SEKTION = re.compile(rf"\b({SEKTIONSORD})\s*({NUM})(?:\s*[-–]\s*({NUM}))?", re.IGNORECASE)
RE_BILAG = re.compile(rf"\b({BILAGSORD})\s+(\d+[a-zA-Z]?|[A-ZÆØÅ])\b", re.IGNORECASE)

RE_MANUELT_NR = re.compile(r"^(\d+(?:\.\d+)+)[\.\)]?\s+\S")   # "4.2 Betaling"
RE_BILAG_DEF = re.compile(rf"^({BILAGSORD})\s+(\d+[a-zA-Z]?|[A-ZÆØÅ])\b", re.IGNORECASE)


def norm_nr(s):
    """Normaliser: '03' -> '3', '4.20.' -> '4.20', 'a' -> 'A'."""
    s = s.strip().rstrip(".")
    dele = []
    for d in s.split("."):
        dele.append(str(int(d)) if d.isdigit() else d.upper())
    return ".".join(dele)


# ---------------------------------------------------------------------------
# Trin 1: Kortlæg definitioner
# ---------------------------------------------------------------------------
def kortlaeg(docs):
    """Returnerer (sektioner_pr_dok, bilag_def).
    sektioner_pr_dok: {doknavn: set af normaliserede numre}
    bilag_def: {(type, nr): definitionssted}  fx ('BILAG','3'): 'Bilag 3 Tidsplan.docx'
    """
    sektioner = {}
    bilag = {}
    for d in docs:
        navn = d.path.name
        numre = set()
        for p_el, label, tekst, _ in d.paras:
            if label and label[:1].isdigit():
                numre.add(norm_nr(label))
            m = RE_MANUELT_NR.match(tekst)
            if m and len(tekst) <= 150:          # ligner en overskrift
                numre.add(norm_nr(m.group(1)))
            m = RE_BILAG_DEF.match(tekst)
            if m and len(tekst) <= 150:
                bilag.setdefault((m.group(1).upper().rstrip("X"), norm_nr(m.group(2))),
                                 f"{navn} (overskrift)")
        sektioner[navn] = numre
        # bilagsidentitet fra filnavnet, fx "Bilag 3 - Tidsplan.docx"
        m = RE_BILAG_DEF.match(d.path.stem.strip())
        if not m:
            m = re.search(rf"\b({BILAGSORD})\s*[_\s-]*(\d+[a-zA-Z]?|[A-ZÆØÅ])\b",
                          d.path.stem, re.IGNORECASE)
        if m:
            bilag[(m.group(1).upper(), norm_nr(m.group(2)))] = navn
    return sektioner, bilag


def bilagstype_match(reftype, deftype):
    """'BILAG' matcher 'BILAG', 'APPENDIKS' matcher 'APPENDIKS'/'APPENDIX'."""
    r, t = reftype.upper(), deftype.upper()
    if r.startswith("APPEND") and t.startswith("APPEND"):
        return True
    return r == t


def find_bilag(bilag_def, reftype, nr):
    """Returnerer (status, sted). 'gyldig' ved præcist match,
    'usikker' ved nær-match (fx Bilag 3 vs kun Bilag 3A defineret)."""
    for (dtype, dnr), sted in bilag_def.items():
        if bilagstype_match(reftype, dtype) and dnr == nr:
            return "gyldig", sted
    for (dtype, dnr), sted in bilag_def.items():
        if bilagstype_match(reftype, dtype) and (dnr.startswith(nr) or nr.startswith(dnr)):
            return "usikker", f"nærmeste match: {dtype.title()} {dnr} ({sted})"
    return "ugyldig", None


# ---------------------------------------------------------------------------
# Trin 2+3: Find og validér henvisninger
# ---------------------------------------------------------------------------
def saetning(tekst, start, slut):
    """Klip den omgivende sætning ud som kontekst."""
    a = max(tekst.rfind(". ", 0, start), tekst.rfind("\n", 0, start)) + 1
    b = tekst.find(". ", slut)
    b = len(tekst) if b == -1 else b + 1
    s = tekst[a:b].strip()
    return (s[:300] + " ...") if len(s) > 300 else s


def er_definition(tekst, m):
    """Spring selve definitionen over (overskriften 'Bilag 3 - Tidsplan'
    eller '4.2 Betaling' er ikke en henvisning)."""
    return m.start() == 0 and len(tekst) <= 150


def validér(docs):
    sektioner, bilag_def = kortlaeg(docs)
    fund = []

    def tjek_sektion(nr, doknavn):
        nr = norm_nr(nr)
        if nr in sektioner[doknavn]:
            return "gyldig", f"afsnit {nr} findes i {doknavn}"
        andre = [n for n, s in sektioner.items() if nr in s and n != doknavn]
        if andre:
            return "usikker", (f"afsnit {nr} findes ikke i {doknavn}, "
                               f"men i {', '.join(andre)} - tjek om henvisningen "
                               f"burde angive dokumentet")
        return "ugyldig", f"afsnit {nr} findes ikke i dokumentsamlingen"

    for d in docs:
        navn = d.path.name
        for p_el, _, tekst, _ in d.paras:
            if not tekst:
                continue
            optaget = []   # tegnpositioner der allerede er matchet

            def ledig(m):
                return not any(a < m.end() and m.start() < b for a, b in optaget)

            # 1) kombineret: "Bilag 3, punkt 2.1"
            for m in RE_BILAG_SEKTION.finditer(tekst):
                if er_definition(tekst, m):
                    continue
                optaget.append((m.start(), m.end()))
                btype, bnr = m.group(1), norm_nr(m.group(2))
                snr = norm_nr(m.group(4))
                bstatus, sted = find_bilag(bilag_def, btype, bnr)
                if bstatus == "ugyldig":
                    status, forkl = "ugyldig", f"{btype.title()} {bnr} findes ikke i samlingen"
                elif bstatus == "usikker":
                    status, forkl = "usikker", sted
                else:
                    if sted.endswith("(overskrift)") :
                        status, forkl = "usikker", (f"{btype.title()} {bnr} er kun nævnt som "
                                                    f"overskrift i {sted.split(' (')[0]} - "
                                                    f"selve bilagsdokumentet er ikke med i samlingen")
                    elif snr in sektioner.get(sted, set()):
                        status, forkl = "gyldig", f"afsnit {snr} findes i {sted}"
                    else:
                        status, forkl = "ugyldig", f"afsnit {snr} findes ikke i {sted}"
                fund.append((status, navn, m.group(0), saetning(tekst, m.start(), m.end()), forkl))

            # 2) bilagshenvisninger: "se Bilag 3"
            for m in RE_BILAG.finditer(tekst):
                if not ledig(m) or er_definition(tekst, m):
                    continue
                optaget.append((m.start(), m.end()))
                btype, bnr = m.group(1), norm_nr(m.group(2))
                status, sted = find_bilag(bilag_def, btype, bnr)
                if status == "gyldig":
                    forkl = f"defineret: {sted}"
                    if sted.endswith("(overskrift)"):
                        status = "usikker"
                        forkl = (f"kun nævnt som overskrift i {sted.split(' (')[0]} - "
                                 f"bilagsdokumentet selv er ikke med i samlingen")
                elif status == "usikker":
                    forkl = sted
                else:
                    forkl = f"{btype.title()} {bnr} findes ikke i dokumentsamlingen"
                fund.append((status, navn, m.group(0), saetning(tekst, m.start(), m.end()), forkl))

            # 3) sektionshenvisninger: "jf. afsnit 4.2" / "punkt 2.1-2.3"
            for m in RE_SEKTION.finditer(tekst):
                if not ledig(m) or er_definition(tekst, m):
                    continue
                optaget.append((m.start(), m.end()))
                status1, forkl = tjek_sektion(m.group(2), navn)
                if m.group(3):                       # interval, fx 2.1-2.3
                    status2, forkl2 = tjek_sektion(m.group(3), navn)
                    par = sorted([status1, status2], key=["gyldig", "usikker", "ugyldig"].index)
                    if par[0] != par[1]:
                        status1, forkl = "usikker", f"interval delvist gyldigt: {forkl} / {forkl2}"
                    else:
                        status1 = par[0]
                        forkl = f"{forkl}; {forkl2}"
                fund.append((status1, navn, m.group(0), saetning(tekst, m.start(), m.end()), forkl))

    return fund, sektioner, bilag_def


# ---------------------------------------------------------------------------
# Rapport
# ---------------------------------------------------------------------------
def skriv_rapport(fund, sektioner, bilag_def, outfile):
    wb = Workbook()

    ws = wb.active
    ws.title = "Rapport"
    hoved = ["Status", "Dokument", "Henvisning", "Sætning (kontekst)", "Vurdering"]
    bredder = [16, 35, 22, 75, 55]
    for c, (h, b) in enumerate(zip(hoved, bredder), 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True)
        ws.column_dimensions[cell.column_letter].width = b
    orden = {"ugyldig": 0, "usikker": 1, "gyldig": 2}
    r = 1
    for status, dok, ref, ktx, forkl in sorted(fund, key=lambda f: orden[f[0]]):
        r += 1
        ws.cell(row=r, column=1, value=SYMBOL[status]).fill = FILLS[status]
        for c, v in enumerate((dok, ref, ktx, forkl), 2):
            ws.cell(row=r, column=c, value=v)
        for c in range(1, 6):
            ws.cell(row=r, column=c).border = THIN
            ws.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")

    ops = wb.create_sheet("Opsummering")
    n = {s: sum(1 for f in fund if f[0] == s) for s in ("gyldig", "ugyldig", "usikker")}
    ops["B1"] = "Opsummering af krydshenvisningstjek"
    ops["B1"].font = Font(bold=True, size=12)
    for i, (s, tekst) in enumerate([("ugyldig", "Ugyldige henvisninger"),
                                    ("usikker", "Usikre henvisninger"),
                                    ("gyldig", "Gyldige henvisninger")], 3):
        ops.cell(row=i, column=2, value=SYMBOL[s]).fill = FILLS[s]
        ops.cell(row=i, column=3, value=tekst)
        ops.cell(row=i, column=4, value=n[s]).font = Font(bold=True)
    ops["B7"] = f"Henvisninger i alt: {len(fund)}"
    ops["B9"] = "Definerede bilag/appendikser i samlingen:"
    ops["B9"].font = Font(bold=True)
    r = 9
    for (btype, bnr), sted in sorted(bilag_def.items()):
        r += 1
        ops.cell(row=r, column=2, value=f"{btype.title()} {bnr}")
        ops.cell(row=r, column=3, value=sted)
    r += 2
    ops.cell(row=r, column=2, value="Dokumenter og antal fundne afsnitsnumre:").font = Font(bold=True)
    for navn, numre in sektioner.items():
        r += 1
        ops.cell(row=r, column=2, value=navn)
        ops.cell(row=r, column=3, value=len(numre))
    for col, b in zip("BCD", (28, 55, 10)):
        ops.column_dimensions[col].width = b

    wb.save(outfile)
    return n


# ---------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description="Validér krydshenvisninger i kontrakt + bilag")
    ap.add_argument("filer", nargs="+", help="Kontrakt og alle bilag (.docx)")
    ap.add_argument("-o", "--output", default="krydstjek.xlsx")
    args = ap.parse_args()

    docs = [ex.Docx(f) for f in args.filer]
    fund, sektioner, bilag_def = validér(docs)
    n = skriv_rapport(fund, sektioner, bilag_def, args.output)
    print(f"Krydstjek gennemført -> {args.output}")
    print(f"  ❌ Ugyldige: {n['ugyldig']}   ⚠️ Usikre: {n['usikker']}   ✅ Gyldige: {n['gyldig']}")


if __name__ == "__main__":
    sys.exit(main())
