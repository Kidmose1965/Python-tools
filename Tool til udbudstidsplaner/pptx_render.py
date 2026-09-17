"""
PPTX-renderer for tidsplan.py - erstatning for draw.js/pptxgenjs.
Tegner den primitiv-liste som byg() i tidsplan.py producerer, med python-pptx.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

W, H = 13.333, 7.5

_ALIGN = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER}
_ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


def _color(hexstr):
    return RGBColor.from_string(hexstr)


def _no_shadow(shp):
    shp.shadow.inherit = False


def _set_transparency(shp, pct):
    if not pct:
        return
    alpha = str(int(round((100 - pct) * 1000)))
    xFill = shp.fill.fore_color._xFill
    a = xFill.makeelement(qn("a:alpha"), {"val": alpha})
    xFill.append(a)


def _set_dash(line, dash):
    if not dash:
        return
    ln = line._get_or_add_ln()
    d = ln.makeelement(qn("a:prstDash"), {"val": dash})
    ln.append(d)


def _shape(slide, kind, p):
    shp = slide.shapes.add_shape(
        kind, Inches(p["x"]), Inches(p["y"]), Inches(max(p["w"], 0.001)), Inches(max(p["h"], 0.001))
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = _color(p["fill"])
    shp.line.fill.background()
    _no_shadow(shp)
    return shp


def _rect(slide, p):
    shp = _shape(slide, MSO_SHAPE.RECTANGLE, p)
    _set_transparency(shp, p.get("transparency", 0))
    return shp


def _diamond(slide, p):
    return _shape(slide, MSO_SHAPE.DIAMOND, p)


def _line(slide, p):
    x1, y1 = Inches(p["x"]), Inches(p["y"])
    x2, y2 = Inches(p["x"] + p["w"]), Inches(p["y"] + p["h"])
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    conn.line.color.rgb = _color(p["color"])
    conn.line.width = Pt(p.get("width", 0.75))
    _set_dash(conn.line, p.get("dash"))
    _no_shadow(conn)
    return conn


def _text(slide, p):
    box = slide.shapes.add_textbox(
        Inches(p["x"]), Inches(p["y"]), Inches(max(p["w"], 0.001)), Inches(max(p["h"], 0.001))
    )
    tf = box.text_frame
    tf.word_wrap = bool(p.get("wrap", False))
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = _ANCHOR.get(p.get("valign", "top"), MSO_ANCHOR.TOP)
    tf.auto_size = None

    s = p["s"]
    runs_spec = s if isinstance(s, list) else [(s, {})]

    paragraphs = [[]]
    for text, style in runs_spec:
        for i, part in enumerate(str(text).split("\n")):
            if i > 0:
                paragraphs.append([])
            if part:
                paragraphs[-1].append((part, style))

    align = _ALIGN.get(p.get("align", "left"), PP_ALIGN.LEFT)
    first = True
    for para_runs in paragraphs:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = align
        for text, style in para_runs:
            run = para.add_run()
            run.text = text
            font = run.font
            font.size = Pt(style.get("size", p.get("size", 10)))
            font.bold = style.get("bold", p.get("bold", False))
            font.italic = style.get("italic", p.get("italic", False))
            font.name = "Calibri"
            font.color.rgb = _color(style.get("color", p.get("color", "000000")))
    return box


_HANDLERS = {
    "rect": _rect,
    "text": _text,
    "diamond": _diamond,
    "line": _line,
}


def render(prims, out_path):
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    for p in prims:
        handler = _HANDLERS.get(p["t"])
        if handler is None:
            raise ValueError(f"Ukendt primitiv: {p['t']}")
        handler(slide, p)
    prs.save(out_path)
