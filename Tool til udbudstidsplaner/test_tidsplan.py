#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Automatiserede tests for tidsplan.py (dato-/varighedsparsing, farvesemantik,
indlaesning og aksebygning). Koeres med:  pytest test_tidsplan.py -v
(eller:  python -m unittest test_tidsplan -v)
Kraever kun det, der allerede er installeret (openpyxl + stdlib unittest,
evt. pytest som testrunner - se requirements.txt).
"""
import datetime as dt
import io
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.styles.colors import Color

import tidsplan as tp

REFERENCEFIL = Path(__file__).resolve().parent / "Testmateriale og baggrund" / "Excel udtræk tidsplan.xlsx"


def _workbook(rows):
    """Bygger et Excel-ark i hukommelsen der efterligner MS Project-udtraekket
    (header i raekke 2, data fra raekke 3). Hver row er en dict med noeglerne:
    navn, indrykket (bool), fed (bool), varighed, start, slut, fill (hex, valgfri),
    tema_farve (Color, valgfri - bruges til at simulere tema-/indekserede farver)."""
    wb = Workbook()
    ws = wb.active
    ws.cell(2, 2, "Aktivitet"); ws.cell(2, 3, "Længde"); ws.cell(2, 4, "Start"); ws.cell(2, 5, "Slut")
    for i, row in enumerate(rows, start=3):
        navn = ("  " + row["navn"]) if row.get("indrykket") else row["navn"]
        c = ws.cell(i, 2, navn)
        if row.get("fed"):
            c.font = Font(bold=True)
        if row.get("fill"):
            c.fill = PatternFill(start_color=row["fill"], end_color=row["fill"], fill_type="solid")
        if row.get("tema_farve") is not None:
            c.fill = PatternFill(fill_type="solid", fgColor=row["tema_farve"])
        ws.cell(i, 3, row.get("varighed", "5 dage"))
        ws.cell(i, 4, row.get("start", ""))
        ws.cell(i, 5, row.get("slut", ""))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _workbook_layout(kolonner, rows, header_row=2, titel=None, gentag_foer=None):
    """Som _workbook(), men med fri overskriftsraekke og kolonnerakkefoelge.
    kolonner: liste af (overskriftstekst, felt) i kolonnerakkefoelge fra kolonne A,
    hvor felt er en af id/navn/varighed/start/slut/pred. titel: tekst i A1.
    gentag_foer: indeks i rows, hvor overskriftsraekken gentages (sideskift)."""
    wb = Workbook()
    ws = wb.active
    if titel:
        ws.cell(1, 1, titel)

    def skriv_overskrift(r):
        for c, (tekst, _) in enumerate(kolonner, start=1):
            ws.cell(r, c, tekst)

    skriv_overskrift(header_row)
    r = header_row + 1
    for i, row in enumerate(rows):
        if gentag_foer == i:
            skriv_overskrift(r)
            r += 1
        for c, (_, felt) in enumerate(kolonner, start=1):
            if felt == "navn":
                cell = ws.cell(r, c, ("  " + row["navn"]) if row.get("indrykket") else row["navn"])
                if row.get("fed"):
                    cell.font = Font(bold=True)
            elif felt == "id":
                ws.cell(r, c, i + 1)
            elif felt == "pred":
                ws.cell(r, c, "")
            else:
                ws.cell(r, c, row.get(felt, "5 days" if felt == "varighed" else ""))
        r += 1
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _msp_raekker():
    """Fase + aktivitet + milepael med MS Project-datoer (ugedag foran)."""
    return [
        dict(navn="Fase 1", fed=True, varighed="4 days", start="Tue 10-11-26", slut="Fri 13-11-26"),
        dict(navn="Aktivitet A", indrykket=True, varighed="3 days", start="Tue 10-11-26", slut="Thu 12-11-26"),
        dict(navn="Milepæl: Aflevering", indrykket=True, varighed="0 days", start="Fri 13-11-26", slut="Fri 13-11-26"),
    ]


MSP_KOLONNER = [("ID", "id"), ("Task Name", "navn"), ("Duration", "varighed"),
                ("Start", "start"), ("Finish", "slut"), ("Predecessors", "pred")]


class TestParseDato(unittest.TestCase):
    def test_date(self):
        self.assertEqual(tp.parse_dato(dt.date(2026, 9, 10)), dt.date(2026, 9, 10))

    def test_datetime(self):
        self.assertEqual(tp.parse_dato(dt.datetime(2026, 9, 10, 13, 30)), dt.date(2026, 9, 10))

    def test_iso_streng(self):
        self.assertEqual(tp.parse_dato("2026-09-10"), dt.date(2026, 9, 10))

    def test_dansk_streng(self):
        self.assertEqual(tp.parse_dato("10-09-2026"), dt.date(2026, 9, 10))

    def test_dansk_streng_2cifret_aar(self):
        self.assertEqual(tp.parse_dato("10-09-26"), dt.date(2026, 9, 10))

    def test_ms_project_ugedag_praefiks(self):
        self.assertEqual(tp.parse_dato("Tue 07-10-25"), dt.date(2025, 10, 7))

    def test_ms_project_ugedag_praefiks_4cifret_aar(self):
        self.assertEqual(tp.parse_dato("Mon 15-09-2026"), dt.date(2026, 9, 15))

    def test_ugyldig_tekst(self):
        with self.assertRaises(ValueError):
            tp.parse_dato("ikke en dato")

    def test_ugyldig_kalenderdato(self):
        with self.assertRaises(ValueError):
            tp.parse_dato("31-02-2026")   # 31. februar findes ikke

    def test_none(self):
        with self.assertRaises(ValueError):
            tp.parse_dato(None)


class TestParseVarighed(unittest.TestCase):
    def test_0_days(self):
        self.assertEqual(tp.parse_varighed("0 days"), 0)

    def test_1_day(self):
        self.assertEqual(tp.parse_varighed("1 day"), 1)

    def test_5_days(self):
        self.assertEqual(tp.parse_varighed("5 days"), 5)

    def test_0_dage(self):
        self.assertEqual(tp.parse_varighed("0 dage"), 0)

    def test_1_dag(self):
        self.assertEqual(tp.parse_varighed("1 dag"), 1)

    def test_5_dage(self):
        self.assertEqual(tp.parse_varighed("5 dage"), 5)

    def test_numerisk_int(self):
        self.assertEqual(tp.parse_varighed(5), 5)

    def test_numerisk_float(self):
        self.assertEqual(tp.parse_varighed(5.0), 5)

    def test_ugyldig_varighed(self):
        with self.assertRaises(ValueError):
            tp.parse_varighed("ukendt")

    def test_ugyldig_varighed_none(self):
        with self.assertRaises(ValueError):
            tp.parse_varighed(None)

    def test_estimat_tegn(self):
        self.assertEqual(tp.parse_varighed("5 days?"), 5)

    def test_decimal_komma_rundes_op_til_1(self):
        self.assertEqual(tp.parse_varighed("0,5 dage"), 1)

    def test_decimal_punktum(self):
        self.assertEqual(tp.parse_varighed("2.5 days"), 3)

    def test_uger(self):
        self.assertEqual(tp.parse_varighed("2 wks"), 10)
        self.assertEqual(tp.parse_varighed("2 uger"), 10)
        self.assertEqual(tp.parse_varighed("1 week"), 5)

    def test_timer(self):
        self.assertEqual(tp.parse_varighed("4 hrs"), 1)     # 0,5 dag -> mindst 1
        self.assertEqual(tp.parse_varighed("16 timer"), 2)

    def test_maaneder(self):
        self.assertEqual(tp.parse_varighed("1 mon"), 20)
        self.assertEqual(tp.parse_varighed("2 måneder"), 40)

    def test_minutter(self):
        self.assertEqual(tp.parse_varighed("480 mins"), 1)
        self.assertEqual(tp.parse_varighed("960 min"), 2)

    def test_elapsed_prefix_ignoreres(self):
        self.assertEqual(tp.parse_varighed("10 edays"), 10)
        self.assertEqual(tp.parse_varighed("2 ewks"), 10)

    def test_ukendt_enhed_er_dage(self):
        self.assertEqual(tp.parse_varighed("3 xyz"), 3)
        self.assertEqual(tp.parse_varighed("3"), 3)

    def test_nul_er_altid_0(self):
        self.assertEqual(tp.parse_varighed("0 hrs"), 0)
        self.assertEqual(tp.parse_varighed("0 days?"), 0)
        self.assertEqual(tp.parse_varighed(0), 0)

    def test_positiv_numerisk_bliver_aldrig_0(self):
        self.assertEqual(tp.parse_varighed(0.4), 1)

    def test_tekst_uden_tal_fejler(self):
        for v in ("Duration", "ukendt", ""):
            with self.subTest(v=v), self.assertRaises(ValueError):
                tp.parse_varighed(v)


class TestFarveSemantik(unittest.TestCase):
    """Farven i kolonne B (Task Name) er semantisk - se laes_plan()/_celle_farve()."""

    def _farve(self, fill=None, tema_farve=None, varighed="5 dage"):
        rows = [dict(navn="Fase 1", fed=True, indrykket=False, varighed=varighed,
                     start="01-09-2026", slut="05-09-2026", fill=fill, tema_farve=tema_farve)]
        lanes, advarsler = tp.laes_plan(_workbook(rows))
        return lanes[0]["tasks"][0], advarsler

    def test_rod_spidsbelastning(self):
        task, _ = self._farve(fill="FF0000")
        self.assertEqual(task[3], tp.RED)

    def test_orange_belastning(self):
        task, _ = self._farve(fill="FFC000")
        self.assertEqual(task[3], tp.ORANGE)

    def test_groen_ekstern(self):
        task, _ = self._farve(fill="92D050")
        self.assertEqual(task[3], tp.GREEN)

    def test_lilla_godkendelse(self):
        task, _ = self._farve(fill="7030A0")
        self.assertEqual(task[3], tp.PURPLE)

    def test_blaa_officiel_milepael(self):
        task, _ = self._farve(fill="00B0F0", varighed="0 dage")
        self.assertIsNone(task[2])
        self.assertEqual(task[3], tp.MS_BLUE)

    def test_orange_moede(self):
        task, _ = self._farve(fill="ED7D31")
        self.assertEqual(task[3], tp.MOEDE)

    def test_hvid_neutral(self):
        task, _ = self._farve(fill="FFFFFF")
        self.assertEqual(task[3], tp.NEUTRAL)

    def test_ukendt_farve_neutral(self):
        task, _ = self._farve(fill=None)
        self.assertEqual(task[3], tp.NEUTRAL)

    def test_tema_farve_neutral_med_advarsel(self):
        task, advarsler = self._farve(tema_farve=Color(theme=4, tint=0.0))
        self.assertEqual(task[3], tp.NEUTRAL)
        self.assertTrue(advarsler)
        self.assertTrue(any("tema" in a.lower() for a in advarsler))

    def test_indekseret_farve_neutral_med_advarsel(self):
        task, advarsler = self._farve(tema_farve=Color(indexed=22))
        self.assertEqual(task[3], tp.NEUTRAL)
        self.assertTrue(advarsler)


class TestMoeder(unittest.TestCase):
    def test_moede_markoer_ved_0_dage(self):
        rows = [dict(navn="Fase 1", fed=True, indrykket=False, varighed="0 dage",
                     start="01-09-2026", slut="01-09-2026", fill="ED7D31")]
        lanes, _ = tp.laes_plan(_workbook(rows))
        task = lanes[0]["tasks"][0]
        self.assertIsNone(task[2])
        self.assertEqual(task[3], tp.MOEDE)
        prims, _ = tp.byg(lanes, dt.date(2026, 9, 1))
        self.assertTrue(any(p["t"] == "diamond" and p["fill"] == tp.MOEDE for p in prims))

    def test_moede_bjaelke_ved_over_0_dage(self):
        rows = [dict(navn="Fase 1", fed=True, indrykket=False, varighed="5 dage",
                     start="01-09-2026", slut="05-09-2026", fill="ED7D31")]
        lanes, _ = tp.laes_plan(_workbook(rows))
        task = lanes[0]["tasks"][0]
        self.assertIsNotNone(task[2])
        self.assertEqual(task[3], tp.MOEDE)
        prims, _ = tp.byg(lanes, dt.date(2026, 9, 1))
        self.assertTrue(any(p["t"] == "rect" and p["fill"] == tp.MOEDE for p in prims))

    def test_moede_paavirker_ikke_belastningsbaand(self):
        # H_LOAD er den faste hoejde paa belastningsbaandets ugerektangler i byg() -
        # en moede-farvet rect med den hoejde ville betyde at moedet fejlagtigt har
        # faerget belastningsbaandet.
        H_LOAD = 0.13
        rows = [
            dict(navn="Fase 1", fed=True, indrykket=False, varighed="5 dage",
                 start="01-09-2026", slut="05-09-2026"),
            dict(navn="Møde", indrykket=True, varighed="5 dage",
                 start="01-09-2026", slut="05-09-2026", fill="ED7D31"),
        ]
        lanes, _ = tp.laes_plan(_workbook(rows))
        prims, _ = tp.byg(lanes, dt.date(2026, 9, 1))
        belastningsbaand_moede = [p for p in prims if p["t"] == "rect" and p["fill"] == tp.MOEDE
                                  and abs(p["h"] - H_LOAD) < 1e-6]
        self.assertEqual(belastningsbaand_moede, [])

    def test_moede_ikke_i_niveau(self):
        self.assertNotIn(tp.MOEDE, tp.NIVEAU)


class TestMilepaele(unittest.TestCase):
    def test_milepael_praefiks_0_dage_blaa(self):
        rows = [dict(navn="Milepæl: Aflevering", fed=True, indrykket=False, varighed="0 dage",
                     start="10-09-2026", slut="10-09-2026")]
        lanes, _ = tp.laes_plan(_workbook(rows))
        task = lanes[0]["tasks"][0]
        self.assertEqual(task[0], "Aflevering")   # praefiks fjernet
        self.assertIsNone(task[2])
        self.assertEqual(task[3], tp.MS_BLUE)

    def test_00B0F0_0_dage_blaa(self):
        rows = [dict(navn="Aflevering", fed=True, indrykket=False, varighed="0 dage",
                     start="10-09-2026", slut="10-09-2026", fill="00B0F0")]
        lanes, _ = tp.laes_plan(_workbook(rows))
        task = lanes[0]["tasks"][0]
        self.assertEqual(task[3], tp.MS_BLUE)

    def test_milepael_praefiks_over_0_dage_ikke_blaa(self):
        rows = [dict(navn="Milepæl: Godkendelsesperiode", fed=True, indrykket=False, varighed="5 dage",
                     start="01-09-2026", slut="05-09-2026")]
        lanes, _ = tp.laes_plan(_workbook(rows))
        task = lanes[0]["tasks"][0]
        self.assertEqual(task[0], "Godkendelsesperiode")  # praefiks fjernet
        self.assertIsNotNone(task[2])                     # almindelig bjaelke, ikke markoer
        self.assertNotEqual(task[3], tp.MS_BLUE)

    def test_almindelig_0_dages_graa(self):
        rows = [dict(navn="Modtagelse", fed=True, indrykket=False, varighed="0 dage",
                     start="10-09-2026", slut="10-09-2026")]
        lanes, _ = tp.laes_plan(_workbook(rows))
        task = lanes[0]["tasks"][0]
        self.assertIsNone(task[2])
        self.assertEqual(task[3], tp.MS_GREY)


class TestLaesPlanOgByg(unittest.TestCase):
    def _normale_raekker(self):
        return [
            dict(navn="Fase 1", fed=True, indrykket=False, varighed="0 dage",
                 start="01-09-2026", slut="01-09-2026"),
            dict(navn="Aktivitet A", indrykket=True, varighed="5 dage",
                 start="01-09-2026", slut="05-09-2026"),
            dict(navn="Milepæl: Aflevering", indrykket=True, varighed="0 dage",
                 start="10-09-2026", slut="10-09-2026"),
        ]

    def test_normal_aktiv_tidsplan(self):
        lanes, advarsler = tp.laes_plan(_workbook(self._normale_raekker()))
        self.assertEqual(len(lanes), 1)
        self.assertEqual(len(lanes[0]["tasks"]), 2)
        self.assertEqual(advarsler, [])
        prims, byg_advarsler = tp.byg(lanes, dt.date(2026, 9, 3))
        self.assertTrue(any(p["t"] == "rect" for p in prims))
        self.assertTrue(any(p["t"] == "diamond" for p in prims))

    def test_statusdato_efter_plan_afslutning(self):
        lanes, _ = tp.laes_plan(_workbook(self._normale_raekker()))
        # Planens sidste dato er 10-09-2026; statusdato er sat langt efter.
        prims, advarsler = tp.byg(lanes, dt.date(2027, 6, 1))
        self.assertTrue(len(prims) > 0)

    def test_byg_uden_aktiviteter_fejler_klart(self):
        with self.assertRaises(ValueError):
            tp.byg([], dt.date(2026, 9, 10))

    def test_manglende_varighed(self):
        rows = [dict(navn="Fase 1", fed=True, indrykket=False, varighed="",
                     start="01-09-2026", slut="05-09-2026")]
        with self.assertRaises(ValueError) as ctx:
            tp.laes_plan(_workbook(rows))
        self.assertIn("række 3", str(ctx.exception).lower())

    def test_ugyldig_varighed_i_plan(self):
        rows = [dict(navn="Fase 1", fed=True, indrykket=False, varighed="ukendt",
                     start="01-09-2026", slut="05-09-2026")]
        with self.assertRaises(ValueError) as ctx:
            tp.laes_plan(_workbook(rows))
        self.assertIn("række 3", str(ctx.exception).lower())

    def test_manglende_startdato(self):
        rows = [dict(navn="Fase 1", fed=True, indrykket=False, varighed="5 dage",
                     start="", slut="05-09-2026")]
        with self.assertRaises(ValueError) as ctx:
            tp.laes_plan(_workbook(rows))
        self.assertIn("række 3", str(ctx.exception).lower())

    def test_ugyldig_startdato(self):
        rows = [dict(navn="Fase 1", fed=True, indrykket=False, varighed="5 dage",
                     start="ikke en dato", slut="05-09-2026")]
        with self.assertRaises(ValueError) as ctx:
            tp.laes_plan(_workbook(rows))
        self.assertIn("række 3", str(ctx.exception).lower())

    def test_manglende_slutdato_almindelig_aktivitet(self):
        rows = [dict(navn="Fase 1", fed=True, indrykket=False, varighed="5 dage",
                     start="01-09-2026", slut="")]
        with self.assertRaises(ValueError) as ctx:
            tp.laes_plan(_workbook(rows))
        self.assertIn("række 3", str(ctx.exception).lower())

    def test_slutdato_foer_startdato(self):
        rows = [dict(navn="Fase 1", fed=True, indrykket=False, varighed="5 dage",
                     start="10-09-2026", slut="01-09-2026")]
        with self.assertRaises(ValueError) as ctx:
            tp.laes_plan(_workbook(rows))
        self.assertIn("slutdato", str(ctx.exception).lower())

    def test_indrykket_uden_fase(self):
        rows = [dict(navn="Aktivitet uden fase", indrykket=True, varighed="5 dage",
                     start="01-09-2026", slut="05-09-2026")]
        with self.assertRaises(ValueError) as ctx:
            tp.laes_plan(_workbook(rows))
        self.assertIn("række 3", str(ctx.exception).lower())

    def test_milepael_uden_gyldig_slutdato_fejler_ikke(self):
        # Slutdato er uden betydning for en milepael og skal ikke blokere indlaesningen.
        rows = [dict(navn="Fase 1", fed=True, indrykket=False, varighed="0 dage",
                     start="01-09-2026", slut="")]
        lanes, _ = tp.laes_plan(_workbook(rows))
        self.assertEqual(len(lanes), 1)

    def test_blank_raekke_ignoreres(self):
        rows = [
            dict(navn="Fase 1", fed=True, indrykket=False, varighed="5 dage",
                 start="01-09-2026", slut="05-09-2026"),
            dict(navn="", varighed="", start="", slut=""),
            dict(navn="Aktivitet A", indrykket=True, varighed="5 dage",
                 start="01-09-2026", slut="05-09-2026"),
        ]
        lanes, _ = tp.laes_plan(_workbook(rows))
        self.assertEqual(len(lanes[0]["tasks"]), 1)

    def test_manglende_task_name_med_data_fejler(self):
        rows = [
            dict(navn="Fase 1", fed=True, indrykket=False, varighed="5 dage",
                 start="01-09-2026", slut="05-09-2026"),
            dict(navn="", varighed="5 dage", start="01-09-2026", slut="05-09-2026"),
        ]
        with self.assertRaises(ValueError) as ctx:
            tp.laes_plan(_workbook(rows))
        self.assertIn("task name", str(ctx.exception).lower())


class TestOverskriftsgenkendelse(unittest.TestCase):
    """Automatisk fund af overskriftsraekken og opslag af kolonner efter navn."""

    def _tjek_msp_plan(self, buf):
        lanes, advarsler = tp.laes_plan(buf)
        self.assertEqual(advarsler, [])
        self.assertEqual(len(lanes), 1)
        self.assertEqual(lanes[0]["name"], "Fase 1")
        self.assertEqual([t[0] for t in lanes[0]["tasks"]], ["Aktivitet A", "Aflevering"])
        self.assertEqual(lanes[0]["tasks"][0][1], dt.date(2026, 11, 10))
        self.assertEqual(lanes[0]["tasks"][0][2], dt.date(2026, 11, 12))
        return lanes

    def test_overskrift_i_raekke_4_med_titel(self):
        # Regression: "Ugyldig varighed: 'Duration'" naar overskriften ikke stod i raekke 2.
        buf = _workbook_layout(MSP_KOLONNER, _msp_raekker(), header_row=4, titel="Udbud - tidsplan")
        self._tjek_msp_plan(buf)

    def test_overskrift_i_raekke_1(self):
        self._tjek_msp_plan(_workbook_layout(MSP_KOLONNER, _msp_raekker(), header_row=1))

    def test_anden_kolonnerakkefoelge(self):
        kolonner = [("Start", "start"), ("Finish", "slut"), ("Task Name", "navn"),
                    ("Duration", "varighed"), ("ID", "id")]
        self._tjek_msp_plan(_workbook_layout(kolonner, _msp_raekker(), header_row=3))

    def test_danske_overskrifter(self):
        kolonner = [("Opgavenavn", "navn"), ("Varighed", "varighed"), ("Start", "start"), ("Slut", "slut")]
        self._tjek_msp_plan(_workbook_layout(kolonner, _msp_raekker(), header_row=2))

    def test_store_bogstaver_og_mellemrum_ligegyldige(self):
        kolonner = [("  TASK NAME ", "navn"), ("duration", "varighed"),
                    (" Start Date", "start"), ("FINISH DATE", "slut")]
        self._tjek_msp_plan(_workbook_layout(kolonner, _msp_raekker(), header_row=2))

    def test_gentaget_overskrift_springes_over(self):
        buf = _workbook_layout(MSP_KOLONNER, _msp_raekker(), header_row=4, titel="Titel", gentag_foer=2)
        lanes = self._tjek_msp_plan(buf)
        self.assertEqual(len(lanes[0]["tasks"]), 2)

    def test_find_overskrift_returnerer_raekke_og_kolonner(self):
        buf = _workbook_layout(MSP_KOLONNER, _msp_raekker(), header_row=4, titel="Titel")
        raekke, kol = tp.find_overskrift(load_workbook(buf).active)
        self.assertEqual(raekke, 4)
        self.assertEqual(kol, dict(navn=2, varighed=3, start=4, slut=5))

    def test_find_overskrift_respekterer_maks(self):
        buf = _workbook_layout(MSP_KOLONNER, _msp_raekker(), header_row=4)
        self.assertIsNone(tp.find_overskrift(load_workbook(buf).active, maks=3))

    def test_fejlbesked_angiver_faktisk_kolonne_og_raekke(self):
        kolonner = [("ID", "id"), ("Task Name", "navn"), ("Predecessors", "pred"),
                    ("Duration", "varighed"), ("Start", "start"), ("Finish", "slut")]
        rows = _msp_raekker()
        rows[1]["varighed"] = "ukendt"          # overskrift i raekke 4 -> fase 5, aktivitet A 6
        with self.assertRaises(ValueError) as ctx:
            tp.laes_plan(_workbook_layout(kolonner, rows, header_row=4, titel="Titel"))
        besked = str(ctx.exception)
        self.assertIn("kolonne D (Duration)", besked)
        self.assertIn("række 6", besked)

    def test_fejlbesked_start_og_finish_kolonner(self):
        rows = _msp_raekker()
        rows[0]["start"] = "ikke en dato"
        with self.assertRaises(ValueError) as ctx:
            tp.laes_plan(_workbook_layout(MSP_KOLONNER, rows, header_row=4))
        self.assertIn("kolonne D (Start)", str(ctx.exception))
        rows = _msp_raekker()
        rows[0]["slut"] = ""
        with self.assertRaises(ValueError) as ctx:
            tp.laes_plan(_workbook_layout(MSP_KOLONNER, rows, header_row=4))
        self.assertIn("kolonne E (Finish)", str(ctx.exception))

    def test_manglende_task_name_angiver_kolonne(self):
        rows = _msp_raekker()
        rows[1]["navn"] = ""
        with self.assertRaises(ValueError) as ctx:
            tp.laes_plan(_workbook_layout(MSP_KOLONNER, rows, header_row=4))
        self.assertIn("Task Name (kolonne B)", str(ctx.exception))

    def test_indrykket_faseindeling_som_i_movia_udtraek(self):
        # Hele hierarkiet er indrykket: faser (fed) med 3 mellemrum, aktiviteter med 6.
        # En fed milepael under en fase (6) er en aktivitet. Samlelinjen over faserne (0) og en
        # fritstaaende milepael paa fase-niveau uden aktiviteter (3) er ikke faser og udelades.
        def raekke(navn, indrykning, fed=False, varighed="5 days", start="Tue 10-11-26", slut="Fri 13-11-26"):
            return dict(navn=" " * indrykning + navn, fed=fed, varighed=varighed, start=start, slut=slut)
        rows = [
            raekke("Milepæl: Start", 0, fed=True, varighed="0 days"),
            raekke("Fase A", 3, fed=True),
            raekke("Aktivitet A1", 6),
            raekke("Milepæl: Godkendt", 6, fed=True, varighed="0 days"),
            raekke("Fase B", 3, fed=True),
            raekke("Aktivitet B1", 6),
            raekke("Aktivitet B2", 6, fed=True),
            raekke("Milepæl: Slut", 3, fed=True, varighed="0 days"),
        ]
        lanes, advarsler = tp.laes_plan(_workbook_layout(MSP_KOLONNER, rows, header_row=4, titel="Titel"))
        self.assertEqual([l["name"] for l in lanes], ["Fase A", "Fase B"])
        self.assertEqual([len(l["tasks"]) for l in lanes], [2, 2])
        self.assertEqual([t[0] for t in lanes[0]["tasks"]], ["Aktivitet A1", "Godkendt"])
        self.assertEqual(len(advarsler), 2)   # udeladt samlelinje (række 5) og slutmilepael (række 12)
        self.assertTrue(any("række 5" in a for a in advarsler))
        self.assertTrue(any("række 12" in a for a in advarsler))

    def test_overskrift_uden_finish_fejler(self):
        kolonner = [("Task Name", "navn"), ("Duration", "varighed"), ("Start", "start")]
        with self.assertRaises(ValueError) as ctx:
            tp.laes_plan(_workbook_layout(kolonner, _msp_raekker(), header_row=3))
        self.assertIn("Finish", str(ctx.exception))

    def test_overskrift_uden_start_fejler(self):
        kolonner = [("Task Name", "navn"), ("Duration", "varighed"), ("Finish", "slut")]
        with self.assertRaises(ValueError) as ctx:
            tp.laes_plan(_workbook_layout(kolonner, _msp_raekker(), header_row=3))
        self.assertIn("Start", str(ctx.exception))

    def test_uden_overskrift_bruges_kolonne_b_e_med_advarsel(self):
        wb = Workbook()
        ws = wb.active
        ws.cell(2, 2, "Noget helt andet"); ws.cell(2, 3, "x"); ws.cell(2, 4, "y"); ws.cell(2, 5, "z")
        c = ws.cell(3, 2, "Fase 1"); c.font = Font(bold=True)
        ws.cell(3, 3, "4 days"); ws.cell(3, 4, "Tue 10-11-26"); ws.cell(3, 5, "Fri 13-11-26")
        buf = io.BytesIO(); wb.save(buf); buf.seek(0)
        lanes, advarsler = tp.laes_plan(buf)
        self.assertEqual(len(lanes), 1)
        self.assertEqual(lanes[0]["tasks"][0][1], dt.date(2026, 11, 10))
        self.assertTrue(any("Overskriftsrække" in a for a in advarsler))

    def test_uden_overskrift_og_uden_kolonne_b_e_fejler(self):
        wb = Workbook()
        ws = wb.active
        ws.cell(1, 1, "kun"); ws.cell(1, 2, "to kolonner")
        buf = io.BytesIO(); wb.save(buf); buf.seek(0)
        with self.assertRaises(ValueError) as ctx:
            tp.laes_plan(buf)
        self.assertIn("B-E", str(ctx.exception))


class TestReferenceFil(unittest.TestCase):
    """End-to-end paa det autoritative modeludtraek fra MS Project."""

    def setUp(self):
        if not REFERENCEFIL.exists():
            self.skipTest(f"Referencefilen findes ikke: {REFERENCEFIL}")

    def test_laeser_75_aktiviteter_i_6_faser(self):
        lanes, advarsler = tp.laes_plan(str(REFERENCEFIL))
        self.assertEqual(len(lanes), 6)
        self.assertEqual(sum(len(l["tasks"]) for l in lanes), 75)

    def test_generering_fejler_ikke_og_legend_har_moede(self):
        konfig = tp.Konfig(kunde="Modelkunde", udbud="Udbud med forhandling - 1 runde")
        lanes, laes_adv = tp.laes_plan(str(REFERENCEFIL), konfig)
        prims, byg_adv = tp.byg(lanes, dt.date(2026, 9, 17), konfig)
        self.assertTrue(len(prims) > 0)
        legend_tekster = [p["s"] for p in prims if p["t"] == "text" and isinstance(p["s"], str)]
        self.assertTrue(any("Møde" == t or "Møde" in t for t in legend_tekster))


if __name__ == "__main__":
    unittest.main()
