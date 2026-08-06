#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Extractor - Python-udgave af "Revised Extractor"
=================================================
Udtrækker kommentarer, trackchanges, markerede inputfelter og kravmatrix
direkte fra .docx-filer - uden at Word åbnes.

Kræver kun: openpyxl   (pip install openpyxl)

Brug:
  python extractor.py kommentarer  FIL.docx [FLERE.docx ...] [-o output.xlsx]
  python extractor.py trackchanges FIL.docx [FLERE.docx ...] [-o output.xlsx]
  python extractor.py inputfelter --farve groen FIL.docx [...] [-o output.xlsx]
  python extractor.py inputfelter --farve gul   FIL.docx [...] [-o output.xlsx]
  python extractor.py kravmatrix   KRAVSPEC.docx            [-o output.xlsx]
  python extractor.py kravmatrix-kommentarer KRAVSPEC.docx  [-o output.xlsx]
  python extractor.py styles       FIL.docx        (viser dokumentets typografier)
"""

import argparse
import os
import posixpath
import re
import sys
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
W14 = "{http://schemas.microsoft.com/office/word/2010/wordml}"
W15 = "{http://schemas.microsoft.com/office/word/2012/wordml}"
BAND = PatternFill("solid", start_color="DCE6F1")
THIN = Border(*[Side(style="thin")] * 4)
BLOKERET = PatternFill("solid", start_color="000000")


def q(tag):
    return W + tag


def lang_sti(path):
    r"""Windows kan ikke åbne stier over ca. 260 tegn uden \\?\-præfikset
    (fx dybe OneDrive-mapper). Path.resolve() giver altid en absolut sti,
    så det er trygt at præfikse den her."""
    s = str(path)
    if os.name == "nt" and not s.startswith("\\\\?\\") and len(s) >= 240:
        s = "\\\\?\\" + s
    return s


# ---------------------------------------------------------------------------
# Nummereringsmotor: genskaber Words "ListString" ud fra numbering.xml
# ---------------------------------------------------------------------------
def _fmt_number(n, numfmt):
    if numfmt == "decimal":
        return str(n)
    if numfmt in ("lowerLetter", "upperLetter"):
        s = ""
        while n > 0:
            n, r = divmod(n - 1, 26)
            s = chr(97 + r) + s
        return s.upper() if numfmt == "upperLetter" else s
    if numfmt in ("lowerRoman", "upperRoman"):
        vals = [(1000, "m"), (900, "cm"), (500, "d"), (400, "cd"), (100, "c"),
                (90, "xc"), (50, "l"), (40, "xl"), (10, "x"), (9, "ix"),
                (5, "v"), (4, "iv"), (1, "i")]
        s = ""
        for v, sym in vals:
            while n >= v:
                s += sym
                n -= v
        return s.upper() if numfmt == "upperRoman" else s
    if numfmt == "none":
        return ""
    return None  # bullets m.v. -> ingen numerisk nummerering


class NumberingEngine:
    def __init__(self, numbering_root):
        self.levels = {}   # numId -> {ilvl: (start, numfmt, lvlText)}
        self.counters = {}  # numId -> {ilvl: aktuel værdi}
        if numbering_root is None:
            return
        abstract = {}
        for an in numbering_root.findall(q("abstractNum")):
            aid = an.get(q("abstractNumId"))
            lvls = {}
            for lvl in an.findall(q("lvl")):
                ilvl = int(lvl.get(q("ilvl")))
                start_el = lvl.find(q("start"))
                fmt_el = lvl.find(q("numFmt"))
                txt_el = lvl.find(q("lvlText"))
                lvls[ilvl] = (
                    int(start_el.get(q("val"))) if start_el is not None else 1,
                    fmt_el.get(q("val")) if fmt_el is not None else "decimal",
                    txt_el.get(q("val")) if txt_el is not None else "%1",
                )
            abstract[aid] = lvls
        for num in numbering_root.findall(q("num")):
            nid = num.get(q("numId"))
            ref = num.find(q("abstractNumId"))
            if ref is not None and ref.get(q("val")) in abstract:
                self.levels[nid] = abstract[ref.get(q("val"))]

    def next_label(self, num_id, ilvl):
        """Tæl op og returnér label, fx '3.2.1' - eller None for punkttegn."""
        lvls = self.levels.get(num_id)
        if not lvls or ilvl not in lvls:
            return None
        c = self.counters.setdefault(num_id, {})
        for lv in range(ilvl):                       # init overliggende niveauer
            if lv in lvls and lv not in c:
                c[lv] = lvls[lv][0]
        c[ilvl] = c.get(ilvl, lvls[ilvl][0] - 1) + 1  # tæl dette niveau op
        for lv in list(c):                            # nulstil dybere niveauer
            if lv > ilvl:
                del c[lv]
        text = lvls[ilvl][2]
        for lv in range(ilvl + 1):
            if lv in lvls:
                part = _fmt_number(c.get(lv, lvls[lv][0]), lvls[lv][1])
                if part is None:
                    return None
                text = text.replace(f"%{lv + 1}", part)
        return text.strip()


# ---------------------------------------------------------------------------
# Docx-læser
# ---------------------------------------------------------------------------
class DocxFejl(Exception):
    """Rejses når en .docx-fil ikke kan læses (beskadiget zip, gammelt
    .doc-format, krypteret/adgangskodebeskyttet, eller mangler sin
    hoveddel)."""


PKG_REL = "{http://schemas.openxmlformats.org/package/2006/relationships}"


def _find_hoveddel(z):
    """Finder stien til dokumentets hoveddel i zip-arkivet. Normalt
    "word/document.xml", men fx Word Online kan navngive den
    "word/document2.xml" e.l. Slår derfor rigtigt op via _rels/.rels
    (relationen med Type ".../relationships/officeDocument") og falder kun
    tilbage til et regex-gæt, hvis relations-filen mangler eller er
    uventet."""
    try:
        rels_root = ET.fromstring(z.read("_rels/.rels"))
        for rel in rels_root.findall(PKG_REL + "Relationship"):
            if rel.get("Type", "").endswith("/relationships/officeDocument"):
                target = rel.get("Target", "").lstrip("/")
                if target in z.namelist():
                    return target
    except (KeyError, ET.ParseError):
        pass
    for name in z.namelist():
        if re.fullmatch(r"word/document\d*\.xml", name):
            return name
    return None


class Docx:
    def __init__(self, path):
        # .resolve() retter Windows' korte 8.3-filnavne (fx "03UDBU~1.DOC" fra
        # tkinters filvalgsdialog) tilbage til det fulde filnavn - uden det
        # fejler bilag-genkendelsen i krydstjek, som matcher på filnavnet.
        self.path = Path(path).resolve()
        sti = lang_sti(self.path)
        try:
            with zipfile.ZipFile(sti) as z:
                hoveddel = _find_hoveddel(z)
                if hoveddel is None:
                    raise DocxFejl(
                        f"{self.path.name}: kunne ikke finde dokumentets "
                        "hoveddel (word/document.xml eller lignende) i "
                        "filen. Filen er muligvis ikke en gyldig .docx-fil.")

                def load(name):
                    try:
                        return ET.fromstring(z.read(name))
                    except KeyError:
                        return None

                self.doc = load(hoveddel)
                if self.doc is None:
                    raise DocxFejl(
                        f"{self.path.name}: dokumentets hoveddel ({hoveddel}) "
                        "kunne ikke findes i filen, selvom relationerne "
                        "peger på den. Filen er muligvis beskadiget.")

                # numbering/comments/styles.xml slås op i samme mappe som
                # hoveddelen selv (normalt "word", men følg med hvis
                # hoveddelen skulle ligge et andet sted), med "word/" som
                # fallback for utraditionelle pakker.
                mappe = posixpath.dirname(hoveddel) or "word"

                def load_i_mappe(navn):
                    data = load(f"{mappe}/{navn}")
                    if data is None and mappe != "word":
                        data = load(f"word/{navn}")
                    return data

                self.numbering = NumberingEngine(load_i_mappe("numbering.xml"))
                self.comments_xml = load_i_mappe("comments.xml")
                self.comments_ext_xml = load_i_mappe("commentsExtended.xml")

                # styleId -> (visningsnavn, numPr-fra-style, basedOn)
                self.styles = {}
                self._style_elements = {}
                styles_root = load_i_mappe("styles.xml")
                if styles_root is not None:
                    for st in styles_root.findall(q("style")):
                        sid = st.get(q("styleId"))
                        name_el = st.find(q("name"))
                        numpr = st.find(f"{q('pPr')}/{q('numPr')}")
                        based = st.find(q("basedOn"))
                        self._style_elements[sid] = st
                        self.styles[sid] = (
                            name_el.get(q("val")) if name_el is not None else sid,
                            numpr,
                            based.get(q("val")) if based is not None else None,
                        )
        except zipfile.BadZipFile:
            magic = b""
            try:
                with open(sti, "rb") as fh:
                    magic = fh.read(4)
            except OSError:
                pass
            if magic == b"\xd0\xcf\x11\xe0":
                raise DocxFejl(
                    f"{self.path.name}: dette er en gammel .doc-fil (eller "
                    "en krypteret/adgangskodebeskyttet .docx-fil), ikke en "
                    "moderne .docx-fil. Gem dokumentet som almindelig .docx "
                    "i Word (fjern evt. adgangskode) og prøv igen.") from None
            raise DocxFejl(
                f"{self.path.name}: filen kunne ikke læses som et "
                "zip-arkiv (.docx-filer er zip-arkiver). Filen er "
                "muligvis beskadiget.") from None

        self._scan()

    # -- styles ------------------------------------------------------------
    def style_name(self, style_id):
        return self.styles.get(style_id, (style_id or "Normal",))[0]

    def _style_numpr(self, style_id, depth=0):
        if style_id not in self.styles or depth > 10:
            return None
        _, numpr, based = self.styles[style_id]
        return numpr if numpr is not None else self._style_numpr(based, depth + 1)

    # -- gennemløb af dokumentet i læserækkefølge ---------------------------
    def _scan(self):
        """Én tur gennem brødteksten: nummerér alle afsnit og notér,
        om de står i en tabel."""
        self.paras = []          # [(element, label|None, tekst, in_table)]
        self.index_of = {}       # id(element) -> indeks i self.paras
        self.celle_ref = {}      # id(afsnit) -> "Tabel 4, række 7, kolonne 2"
        self._tabel_nr = 0
        self.raekke_aendring = {}   # id(rækkens første afsnit) -> beskrivelse
        self.raekke_afsnit = set()  # id(afsnit) der ligger i en ændret række
        body = self.doc.find(q("body"))
        if body is None:
            raise DocxFejl(
                f"{self.path.name}: dokumentets hoveddel mangler et "
                "<w:body>-element. Filen er muligvis beskadiget eller ikke "
                "en gyldig .docx-fil.")
        for el, in_tbl in self._iter_block(body, False):
            label = self._para_label(el)
            text = para_text(el, include_del=True)
            if label is None:
                label = self._tekst_label(el, text)
            self.index_of[id(el)] = len(self.paras)
            self.paras.append((el, label, text, in_tbl))

    def _iter_block(self, parent, in_table, ref=None):
        for child in parent:
            if child.tag == q("p"):
                if ref:
                    self.celle_ref[id(child)] = ref
                yield child, in_table
            elif child.tag == q("tbl"):
                self._tabel_nr += 1
                tnr = self._tabel_nr
                for ri, tr in enumerate(child.findall(q("tr")), 1):
                    # Hele rækker, der er slettet eller indsat med registrering
                    # af ændringer, markeres med <w:ins>/<w:del> inde i
                    # <w:trPr> - ikke som tekstændringer. Uden det kan man
                    # ikke skelne "et ord er rettet inde i en række" fra
                    # "hele rækken er væk".
                    mark = None
                    trpr = tr.find(q("trPr"))
                    if trpr is not None:
                        for tag, typ in ((q("del"), "Række slettet"),
                                         (q("ins"), "Række indsat")):
                            el = trpr.find(tag)
                            if el is not None:
                                mark = {"type": typ,
                                        "forfatter": el.get(q("author")) or "",
                                        "placering": f"Tabel {tnr}, række {ri}",
                                        "foerste": None}
                                break
                    for ci, tc in enumerate(tr.findall(q("tc")), 1):
                        for p_el, it in self._iter_block(
                                tc, True, f"Tabel {tnr}, række {ri}, kolonne {ci}"):
                            if mark is not None:
                                if mark["foerste"] is None:
                                    mark["foerste"] = p_el
                                    self.raekke_aendring[id(p_el)] = mark
                                self.raekke_afsnit.add(id(p_el))
                            yield p_el, it
                    if mark is not None:
                        mark["tekst"] = " | ".join(
                            t for t in (para_text(pe, include_del=True)
                                        for tc in tr.findall(q("tc"))
                                        for pe in tc.iter(q("p"))) if t)

    # Overskrift med manuelt indtastet nummer, fx "3.1 - Dyrlæger" eller
    # "6.1.1\tRecept". Kræver at der står tekst efter nummeret.
    _MANUELT_NR = re.compile(r"^\s*(\d+(?:\.\d+)*)\s*(?:[-–—.)\t ]\s*)(\S.*)$")

    def _style_element(self, sid):
        return self._style_elements.get(sid)

    def _er_overskrift(self, p):
        """True hvis afsnittet er en overskrift - enten via de indbyggede
        typografinavne ("heading 1"..."heading 9") eller via outlineLvl,
        som brugerdefinerede overskriftstypografier bruger."""
        if re.match(r"heading [1-9]$", self.para_style_name(p) or "", re.I):
            return True
        ppr = p.find(q("pPr"))
        if ppr is not None:
            if ppr.find(q("outlineLvl")) is not None:
                return True
            ps = ppr.find(q("pStyle"))
            if ps is not None:
                sid = ps.get(q("val"))
                for _ in range(10):          # følg basedOn-kæden
                    if sid not in self.styles:
                        break
                    st_el = self._style_element(sid)
                    if st_el is not None and st_el.find(
                            f"{q('pPr')}/{q('outlineLvl')}") is not None:
                        return True
                    sid = self.styles[sid][2]
        return False

    def _tekst_label(self, p, text):
        """Fallback for dokumenter uden automatisk nummerering: læs nummeret
        ud af selve overskriftsteksten."""
        if not text or not self._er_overskrift(p):
            return None
        m = self._MANUELT_NR.match(text)
        return m.group(1) if m else None

    def _para_label(self, p):
        ppr = p.find(q("pPr"))
        numpr = ppr.find(q("numPr")) if ppr is not None else None
        if numpr is None and ppr is not None:
            pstyle = ppr.find(q("pStyle"))
            if pstyle is not None:
                numpr = self._style_numpr(pstyle.get(q("val")))
        if numpr is None:
            return None
        nid_el = numpr.find(q("numId"))
        ilvl_el = numpr.find(q("ilvl"))
        if nid_el is None:
            return None
        nid = nid_el.get(q("val"))
        ilvl = int(ilvl_el.get(q("val"))) if ilvl_el is not None else 0
        if nid == "0":
            return None
        return self.numbering.next_label(nid, ilvl)

    # -- "nærmeste nummererede overskrift" (som i VBA-værktøjet) ------------
    def section_for(self, p_element):
        i = self.index_of.get(id(p_element))
        if i is None:
            return "Ukendt placering"
        while i >= 0:
            _, label, text, _ = self.paras[i]
            if label and label[:1].isdigit():
                if text.lstrip().startswith(label):
                    return text.strip()
                return f"{label} - {text}".strip(" -")
            i -= 1
        return "Før nummereret sektion"

    def para_style_name(self, p):
        ppr = p.find(q("pPr"))
        if ppr is not None:
            ps = ppr.find(q("pStyle"))
            if ps is not None:
                return self.style_name(ps.get(q("val")))
        return self.style_name(None)


def para_text(p, include_del=False):
    parts = []
    for node in p.iter():
        if node.tag == q("t") and node.text:
            parts.append(node.text)
        elif include_del and node.tag == q("delText") and node.text:
            parts.append(node.text)
        elif node.tag == q("tab"):
            parts.append("\t")
    return "".join(parts).strip()


# ---------------------------------------------------------------------------
# Udtræk 1: Kommentarer
# ---------------------------------------------------------------------------
_DATOFORMATER = (
    "%Y-%m-%dT%H:%M:%S",
    "%Y-%m-%d",
    "%d/%m/%Y %H:%M:%S",
    "%d/%m/%Y",
    "%d-%m-%Y %H:%M:%S",
    "%d-%m-%Y",
)


def _normaliser_dato(raw):
    """Tolker w:date i flere kendte formater og returnerer altid ÅÅÅÅ-MM-DD.
    Word skriver typisk ISO 8601 (med eller uden Z-suffiks), men enkelte
    kommentarer - fx tilføjet af et andet værktøj - kan have et helt andet
    format. Kan værdien ikke tolkes, returneres den rå streng uændret;
    funktionen returnerer aldrig tom streng (for en ikke-tom input) og
    kaster aldrig."""
    if not raw:
        return raw
    s = raw.strip()
    if s.endswith("Z"):
        s = s[:-1]
    for fmt in _DATOFORMATER:
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return raw


def _naeste_ikke_tomme_tekst(docx, a_el, maks=3):
    """Nærmeste efterfølgende ikke-tomme afsnit efter ankeret, højst `maks`
    afsnit frem. Bruges som fallback for punktkommentarer hvor selve
    ankerafsnittet er tomt (fx en tom tabelcelle eller et afstandsafsnit)."""
    i = docx.index_of.get(id(a_el))
    if i is None:
        return ""
    for j in range(i + 1, min(i + 1 + maks, len(docx.paras))):
        txt = para_text(docx.paras[j][0])
        if txt:
            return txt
    return ""


def extract_comments(docx):
    out = []
    if docx.comments_xml is None:
        return out
    meta = {}
    para_to_cid = {}
    for c in docx.comments_xml.findall(q("comment")):
        cid = c.get(q("id"))
        ps = c.findall(q("p"))
        text = "\n".join(filter(None, (para_text(p) for p in ps)))
        meta[cid] = (c.get(q("initials")) or c.get(q("author")) or "",
                     text, _normaliser_dato(c.get(q("date")) or ""))
        if ps:
            pid = ps[-1].get(W14 + "paraId")
            if pid:
                para_to_cid[pid] = cid

    parent_of = {}
    if docx.comments_ext_xml is not None:
        for cex in docx.comments_ext_xml.iter(W15 + "commentEx"):
            pid = cex.get(W15 + "paraId")
            parent_pid = cex.get(W15 + "paraIdParent")
            if pid in para_to_cid and parent_pid in para_to_cid:
                parent_of[para_to_cid[pid]] = para_to_cid[parent_pid]

    # scope-tekst, ankerafsnit OG dokumentposition.
    # Word tildeler IKKE kommentar-id'er i dokumentrækkefølge (nogle id'er er
    # tilfældige 32-bit tal), så positionen skal aflæses af dokumentet selv.
    scope_dele, anchor, anchor_pos, active = {}, {}, {}, set()
    for p_i, (p_el, _, _, _) in enumerate(docx.paras):
        loebende = {cid: [] for cid in active}
        for n_i, node in enumerate(p_el.iter()):
            if node.tag == q("commentRangeStart"):
                cid = node.get(q("id"))
                active.add(cid)
                loebende.setdefault(cid, [])
                scope_dele.setdefault(cid, [])
                anchor.setdefault(cid, p_el)
                anchor_pos.setdefault(cid, (p_i, n_i))
            elif node.tag == q("commentRangeEnd"):
                active.discard(node.get(q("id")))
            elif node.tag == q("t") and node.text and active:
                for cid in active:
                    loebende.setdefault(cid, []).append(node.text)
            elif node.tag == q("commentReference"):
                cid = node.get(q("id"))
                anchor.setdefault(cid, p_el)
                anchor_pos.setdefault(cid, (p_i, n_i))
        # ét afsnit/én tabelcelle = ét stykke, så tekst fra to celler ikke
        # løber sammen ("PakningsstoerrelseTekstF.eks. ...")
        for cid, dele in loebende.items():
            t = "".join(dele).strip()
            if t:
                scope_dele.setdefault(cid, []).append(t)

    SIDST = (10 ** 9, 0)

    def pos(cid):
        return anchor_pos.get(cid, SIDST)

    # Traade: rod foerst, derefter svar sorteret efter dato (id-raekkefoelge
    # er ubrugelig - se ovenfor). Roden bestemmer traadens plads i arket.
    def rod(cid):
        r, hop = cid, 0
        while r in parent_of and hop < 50:
            r = parent_of[r]
            hop += 1
        return r

    svar_til = {}
    for cid in meta:
        r = rod(cid)
        if r != cid:
            svar_til.setdefault(r, []).append(cid)

    raekkefoelge = []
    for r in sorted((c for c in meta if rod(c) == c), key=lambda c: (pos(c), meta[c][2])):
        raekkefoelge.append(r)
        for s_cid in sorted(svar_til.get(r, []), key=lambda c: (meta[c][2], pos(c))):
            raekkefoelge.append(s_cid)
    for cid in meta:                      # sikkerhedsnet
        if cid not in raekkefoelge:
            raekkefoelge.append(cid)

    num_of = {cid: n for n, cid in enumerate(raekkefoelge, 1)}
    traad_nr, t = {}, 0
    for cid in raekkefoelge:
        r = rod(cid)
        if r not in traad_nr:
            t += 1
            traad_nr[r] = t

    for cid in raekkefoelge:
        initials, text, dato = meta[cid]
        a_el = anchor.get(cid)
        sec = docx.section_for(a_el) if a_el is not None else "Ukendt placering"
        markeret = " | ".join(scope_dele.get(cid, [])).strip()
        if not markeret and a_el is not None:
            # punktkommentar: markoeren stod i teksten uden at der var
            # markeret noget. Vis afsnittet den haenger paa, ellers staar
            # kolonnen tom og kommentaren er ulaeselig ude af kontekst.
            # Er ankerafsnittet selv tomt (fx en tom tabelcelle), fald
            # videre til det naermeste efterfoelgende ikke-tomme afsnit.
            afsnit = para_text(a_el) or _naeste_ikke_tomme_tekst(docx, a_el)
            if afsnit:
                markeret = f"(punktkommentar) {afsnit[:300]}"
        r = rod(cid)
        out.append({
            "Dokument": docx.path.name, "Nummer": num_of[cid],
            "Tråd": traad_nr[r],
            "Type": "Ny tråd" if r == cid else "Svar",
            "Svar på": "" if r == cid else f"Nr. {num_of[parent_of[cid]]}",
            "Dato": dato, "Kommentar": text,
            "Markeret tekst": markeret,
            "Initialer": initials, "Nummereret sektion": sec,
        })
    return out


# ---------------------------------------------------------------------------
# Udtræk 1b: Kommentarer fra Excel
# ---------------------------------------------------------------------------
def extract_excel_comments(path):
    path = Path(path).resolve()
    wb = load_workbook(lang_sti(path), read_only=False, data_only=True)
    out = []
    n = 0
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows():
            for cell in row:
                if cell.comment:
                    n += 1
                    out.append({
                        "Dokument": path.name,
                        "Nummer": n,
                        "Kommentar": str(cell.comment.text or "").strip(),
                        "Markeret tekst": str(cell.value) if cell.value is not None else "",
                        "Initialer": cell.comment.author or "",
                        "Nummereret sektion": f"{sheet_name} - {cell.column_letter}{cell.row}",
                    })
    wb.close()
    return out


# ---------------------------------------------------------------------------
# Udtræk 2: Trackchanges
# ---------------------------------------------------------------------------
def _kontekst(p_el):
    """Hele afsnittet med {-slettet-} / {+indsat+} markeret."""
    ud = []

    def gaa(el, tilstand):
        for barn in el:
            t = barn.tag
            ny = tilstand
            if t == q("ins"):
                ny = "ins"
            elif t == q("del"):
                ny = "del"
            elif t == q("pPr"):
                continue          # afsnitsmærke-ændringer er ikke tekst
            if t == q("t") and barn.text:
                ud.append("{+%s+}" % barn.text if tilstand == "ins"
                          else "{-%s-}" % barn.text if tilstand == "del"
                          else barn.text)
            elif t == q("delText") and barn.text:
                ud.append("{-%s-}" % barn.text)
            else:
                gaa(barn, ny)

    gaa(p_el, None)
    return "".join(ud).strip()


def _placering(docx, p_el):
    return docx.celle_ref.get(id(p_el), "brødtekst")


def extract_trackchanges(docx, saml=True):
    out = []
    for p_el, _, _, _ in docx.paras:
        mark = docx.raekke_aendring.get(id(p_el))
        if mark is not None:
            out.append({"Dokument": docx.path.name,
                        "Sektion": docx.section_for(p_el),
                        "Placering": mark["placering"],
                        "Ændring": mark.get("tekst") or "(tom række)",
                        "Type": mark["type"],
                        "Kontekst": "",
                        "Forfatter": mark["forfatter"]})
        if id(p_el) in docx.raekke_afsnit:
            continue          # rækken er rapporteret samlet ovenfor
        sektion = docx.section_for(p_el)
        kontekst = _kontekst(p_el)
        placering = _placering(docx, p_el)
        fund = []
        for node in p_el.iter():          # .iter(), ikke direkte børn - se note
            if node.tag == q("ins"):
                txt = para_text(node)
                if txt:
                    fund.append(("Indsat", txt, node.get(q("author")) or ""))
            elif node.tag == q("del"):
                txt = para_text(node, include_del=True)
                if txt:
                    fund.append(("Slettet", txt, node.get(q("author")) or ""))
        if saml:
            # Word splitter ofte én redigering op i mange <w:ins>/<w:del>
            samlet = []
            for typ, txt, forf in fund:
                if samlet and samlet[-1][0] == typ and samlet[-1][2] == forf:
                    samlet[-1][1] += txt
                else:
                    samlet.append([typ, txt, forf])
            fund = [tuple(x) for x in samlet]
        for typ, txt, forf in fund:
            out.append({"Dokument": docx.path.name,
                        "Sektion": sektion,
                        "Placering": placering,
                        "Ændring": txt, "Type": typ,
                        "Kontekst": kontekst,
                        "Forfatter": forf})
    return out


# ---------------------------------------------------------------------------
# Udtræk 3: Markerede inputfelter (highlight)
# ---------------------------------------------------------------------------
HIGHLIGHT = {"groen": "green", "grøn": "green", "gul": "yellow"}


def extract_highlights(docx, colour):
    out = []
    want = HIGHLIGHT[colour]
    for p_el, _, _, _ in docx.paras:
        current = []
        for r in p_el.iter(q("r")):
            rpr = r.find(q("rPr"))
            hl = rpr.find(q("highlight")) if rpr is not None else None
            if hl is not None and hl.get(q("val")) == want:
                current.append(para_text(r))
            elif current:
                out.append({"Dokument": docx.path.name,
                            "Sektion": docx.section_for(p_el),
                            "Inputfelt": "".join(current)})
                current = []
        if current:
            out.append({"Dokument": docx.path.name,
                        "Sektion": docx.section_for(p_el),
                        "Inputfelt": "".join(current)})
    return out


# ---------------------------------------------------------------------------
# Udtræk 4: Kravmatrix
# ---------------------------------------------------------------------------
def _comments_by_anchor_paragraph(docx):
    """id(afsnit) -> liste af kommentartekster der starter i det afsnit."""
    if docx.comments_xml is None:
        return {}
    meta = {}
    for c in docx.comments_xml.findall(q("comment")):
        cid = c.get(q("id"))
        text = "\n".join(filter(None, (para_text(p) for p in c.findall(q("p")))))
        meta[cid] = text

    anchor = {}
    for p_el, _, _, _ in docx.paras:
        for node in p_el.iter():
            if node.tag in (q("commentRangeStart"), q("commentReference")):
                anchor.setdefault(node.get(q("id")), p_el)

    by_para = {}
    for cid, p_el in anchor.items():
        if cid in meta:
            by_para.setdefault(id(p_el), []).append(meta[cid])
    return by_para


def list_styles(docx):
    seen, ordered = set(), []
    for p_el, _, _, _ in docx.paras:
        nm = docx.para_style_name(p_el)
        if nm not in seen:
            seen.add(nm)
            ordered.append(nm)
    return ordered


def _pick(prompt, styles, allow_all):
    print(prompt)
    extra = "  0: [ALLE]\n" if allow_all else "  0: [INGEN/færdig]\n"
    print(extra + "\n".join(f"  {i}: {s}" for i, s in enumerate(styles, 1)))
    raw = input("Angiv numre adskilt af komma (fx 2,5): ").strip()
    chosen = []
    for tok in raw.split(","):
        tok = tok.strip()
        if tok == "0" and allow_all:
            return "ALLE"
        if tok.isdigit() and 1 <= int(tok) <= len(styles):
            chosen.append(styles[int(tok) - 1])
    return chosen


def extract_kravmatrix(docx, table_styles=None, heading_styles=None, with_comments=False):
    if table_styles is None:
        styles = list_styles(docx)
        table_styles = _pick("\nStyles tilhørende TABEL-elementer (kravene):", styles, True)
        heading_styles = _pick("\nStyles tilhørende OVERSKRIFTER (separatorer):", styles, False) or []
    get_all = table_styles == "ALLE"
    comments_by_para = _comments_by_anchor_paragraph(docx) if with_comments else {}

    rows = []
    body = docx.doc.find(q("body"))
    for child in body:
        if child.tag == q("tbl"):
            for tr in child.findall(q("tr")):
                tcs = tr.findall(q("tc"))
                if not tcs:
                    continue
                p1 = tcs[0].find(q("p"))
                style = docx.para_style_name(p1) if p1 is not None else ""
                if not (get_all or style in table_styles):
                    continue
                c1 = cell_text(docx, tcs[0])
                c2 = cell_text(docx, tcs[1]) if len(tcs) > 1 else ""
                if not (c1 or c2):
                    continue
                if with_comments:
                    kommentarer = [txt for tc in tcs for p in tc.findall(q("p"))
                                   for txt in comments_by_para.get(id(p), [])]
                    rows.append(("krav", c1, c2, "\n".join(kommentarer)))
                else:
                    rows.append(("krav", c1, c2))
        elif child.tag == q("p"):
            style = docx.para_style_name(child)
            if style in (heading_styles or []):
                i = docx.index_of.get(id(child))
                label = docx.paras[i][1] if i is not None else None
                text = para_text(child)
                if text:
                    rows.append(("overskrift", label or "", text, "") if with_comments
                               else ("overskrift", label or "", text))
    return rows


def cell_text(docx, tc):
    parts = []
    for p in tc.findall(q("p")):
        i = docx.index_of.get(id(p))
        label = docx.paras[i][1] if i is not None else None
        txt = para_text(p)
        parts.append(f"{label} {txt}".strip() if label else txt)
    return "\n".join(filter(None, parts)).strip()


# ---------------------------------------------------------------------------
# Excel-output
# ---------------------------------------------------------------------------
def write_simple(records, headers, outfile, widths):
    wb = Workbook()
    # openpyxl stempler ellers created/modified med "nu", hvilket gør to
    # ellers identiske kørsler forskellige på byte-niveau (kun i
    # docProps/core.xml - selve dataarket er upåvirket).
    wb.properties.created = wb.properties.modified = datetime(1970, 1, 1)
    ws = wb.active
    for col, (h, w) in enumerate(zip(headers, widths), 2):
        c = ws.cell(row=1, column=col, value=h)
        c.font = Font(bold=True)
        ws.column_dimensions[c.column_letter].width = w
    INGEN_FUND_FILL = PatternFill("solid", start_color="CCEBF9")
    for r, rec in enumerate(records, 2):
        ingen_fund = rec.get("__ingen_fund__", False)
        for col, h in enumerate(headers, 2):
            c = ws.cell(row=r, column=col, value=rec.get(h, ""))
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if ingen_fund:
                c.fill = INGEN_FUND_FILL
    for row in ws.iter_rows(min_row=1, max_row=max(1, len(records) + 1),
                            min_col=2, max_col=len(headers) + 1):
        for c in row:
            c.border = THIN
    wb.save(outfile)


def write_kravmatrix(rows, outfile, tilbudsgiver="[…]"):
    wb = Workbook()
    ws = wb.active
    ws["B1"] = "Bilag XX, Appendiks XX - Kravmatrix"
    ws["B1"].font = Font(bold=True, size=12)
    ws["B3"] = f"Navn på tilbudsgiver: {tilbudsgiver}"
    headers = [("B5", "Krav nr."), ("C5", "Kravbeskrivelse"), ("D5", "Krav opfyldt"),
               ("G5", "Standard-programmel"), ("H5", "Tilpasning/opsætning"),
               ("I5", "Redegørelse"), ("D6", "J"), ("E6", "N"), ("F6", "D")]
    for ref, val in headers:
        ws[ref] = val
        ws[ref].font = Font(bold=True)
    ws.merge_cells("D5:F5")
    for col, w in zip("BCDEFGHI", (21, 85, 6, 6, 6, 20, 22, 22)):
        ws.column_dimensions[col].width = w

    mindstekrav_raekker = []
    r = 6
    for kind, a, b in rows:
        r += 1
        ws.cell(row=r, column=2, value=a)
        c_krav = ws.cell(row=r, column=3, value=b)
        if kind == "overskrift":
            ws.cell(row=r, column=2).font = Font(bold=True)
            c_krav.font = Font(bold=True)
        elif "[MK]" in b:
            c_krav.font = Font(bold=True)
            mindstekrav_raekker.append(r)
    if r > 6:
        for i in range(7, r + 1):
            for col in range(2, 10):
                c = ws.cell(row=i, column=col)
                c.border = THIN
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if i % 2 == 0:
                    c.fill = BAND
        # Mindstekrav: lås N, D, Standard-programmel, Tilpasning/opsætning og
        # Redegørelse (sort baggrund) - kun J (kolonne D) forbliver hvid/åben.
        for i in mindstekrav_raekker:
            for col in (5, 6, 7, 8, 9):
                ws.cell(row=i, column=col).fill = BLOKERET
    wb.save(outfile)


def write_kravmatrix_kommentarer(rows, outfile):
    wb = Workbook()
    ws = wb.active
    headers = [("B1", "Krav nr."), ("C1", "Kravbeskrivelse"), ("D1", "Kommentarer")]
    for ref, val in headers:
        ws[ref] = val
        ws[ref].font = Font(bold=True)
    for col, w in zip("BCD", (21, 85, 60)):
        ws.column_dimensions[col].width = w

    r = 1
    for kind, a, b, kommentar in rows:
        r += 1
        ws.cell(row=r, column=2, value=a)
        c_krav = ws.cell(row=r, column=3, value=b)
        ws.cell(row=r, column=4, value=kommentar)
        if kind == "overskrift":
            ws.cell(row=r, column=2).font = Font(bold=True)
            c_krav.font = Font(bold=True)
        elif "[MK]" in b:
            # Mindstekrav: kravteksten gøres fed, men kommentar-kolonnen
            # (reviewerens felt) forbliver hvid og åben - lås IKKE den her.
            c_krav.font = Font(bold=True)
    if r > 1:
        for i in range(2, r + 1):
            for col in range(2, 5):
                c = ws.cell(row=i, column=col)
                c.border = THIN
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if i % 2 == 0:
                    c.fill = BAND
    wb.save(outfile)


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description="Udtræk fra Word-kravspecifikationer (.docx)")
    sub = ap.add_subparsers(dest="cmd", required=True)
    for navn in ("kommentarer", "trackchanges", "kravmatrix", "kravmatrix-kommentarer", "styles"):
        p = sub.add_parser(navn)
        p.add_argument("filer", nargs="+")
        p.add_argument("-o", "--output")
    p = sub.add_parser("inputfelter")
    p.add_argument("filer", nargs="+")
    p.add_argument("--farve", choices=["groen", "grøn", "gul"], required=True)
    p.add_argument("-o", "--output")
    args = ap.parse_args()

    docs = [Docx(f) for f in args.filer]

    if args.cmd == "styles":
        for d in docs:
            print(f"\n{d.path.name}:")
            for s in list_styles(d):
                print(f"  - {s}")
        return

    if args.cmd == "kommentarer":
        recs = [r for d in docs for r in (extract_comments(d) or
                [{"Dokument": d.path.name, "Kommentar": "Ingen kommentarer",
                  "__ingen_fund__": True}])]
        out = args.output or "kommentarer.xlsx"
        write_simple(recs, ["Dokument", "Nummer", "Tråd", "Type", "Svar på",
                            "Dato", "Kommentar", "Markeret tekst", "Initialer",
                            "Nummereret sektion"], out,
                     [35, 8, 7, 9, 9, 11, 50, 50, 10, 40])

    elif args.cmd == "trackchanges":
        recs = [r for d in docs for r in (extract_trackchanges(d) or
                [{"Dokument": d.path.name, "Ændring": "Ingen trackchanges"}])]
        forfattere = {r.get("Forfatter") for r in recs if "Forfatter" in r}
        if len(forfattere) == 1:
            navn = next(iter(forfattere))
            if navn:
                antal = sum(1 for r in recs if "Forfatter" in r)
                print(f'ADVARSEL: alle {antal} ændringer er tilskrevet "{navn}". Er '
                      "filen et Word Compare-output, er forfatteren den der kørte "
                      "sammenligningen - ikke den der lavede ændringerne.")
        out = args.output or "trackchanges.xlsx"
        write_simple(recs, ["Dokument", "Sektion", "Placering", "Ændring",
                            "Type", "Kontekst", "Forfatter"],
                     out, [35, 45, 26, 55, 10, 90, 20])

    elif args.cmd == "inputfelter":
        recs = [r for d in docs for r in (extract_highlights(d, args.farve) or
                [{"Dokument": d.path.name, "Inputfelt": "Ingen markerede felter"}])]
        out = args.output or f"inputfelter_{args.farve}.xlsx"
        write_simple(recs, ["Dokument", "Sektion", "Inputfelt"], out, [35, 50, 60])

    elif args.cmd == "kravmatrix":
        rows = []
        for d in docs:
            rows += extract_kravmatrix(d)
        out = args.output or "kravmatrix.xlsx"
        write_kravmatrix(rows, out)

    elif args.cmd == "kravmatrix-kommentarer":
        rows = []
        for d in docs:
            rows += extract_kravmatrix(d, with_comments=True)
        out = args.output or "kravmatrix_kommentarer.xlsx"
        write_kravmatrix_kommentarer(rows, out)

    er_kravmatrix = args.cmd in ("kravmatrix", "kravmatrix-kommentarer")
    print(f"Udtræk gennemført -> {out}  ({len(rows) if er_kravmatrix else len(recs)} rækker)")


if __name__ == "__main__":
    sys.exit(main())
