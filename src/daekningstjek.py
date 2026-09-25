#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Dækningstjek - opdager afsnit i et måldokument, der ALDRIG citeres fra et
sæt kildedokumenter
============================================================================
Modsat krydstjek.py (validerer at HVER ENKELT henvisning der findes i
teksten peger på noget der eksisterer) og intern_tjek.py (samme, for ét
dokument) tjekker dette modul den MODSATTE retning: er der afsnit i
måldokumentet, som INGEN henvisning fra kildedokumenterne nogensinde rammer?

Den slags "stille" fejl kan hverken krydstjek.py eller intern_tjek.py
opdage, fordi de kun ser på henvisninger der rent faktisk står i teksten -
et afsnit der aldrig bliver nævnt giver jo ingen henvisning at validere imod.

Typisk brug: et systematisk evaluerings-/tildelingskriterie-afsnit i
udbudsbetingelserne, der gennemgår et andet dokuments afsnit kapitel for
kapitel (fx "bilag 1A, afsnit 4.1, 4.2 ... 6.7 imødekommes"). Dækningstjek
afslører hvis et afsnit i målet er faldet ud af den gennemgang - fx fordi
en senere tilføjet sektion har forskudt nummereringen et sted i en liste.

Kan køres alene:
    python daekningstjek.py udbudsbetingelser.docx --maal kravspec.docx --fra 4 --til 6 -o daekningstjek.xlsx
eller via knappen "Dækningstjek" i extractor_gui.py.

Kræver: extractor.py og krydstjek.py i samme mappe + openpyxl.
"""

import argparse
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

import extractor as ex
import krydstjek as kt

THIN = Border(*[Side(style="thin")] * 4)
FILL_DAEKKET = PatternFill("solid", start_color="C6E0B4")
FILL_MANGLER = PatternFill("solid", start_color="F8CBAD")


def _sort_key(nr):
    """Numerisk sortering af normaliserede afsnitsnumre ('4.10' efter '4.9',
    ikke mellem '4.1' og '4.2' som en ren tekst-sortering ville give)."""
    ud = []
    for del_ in nr.split("."):
        m = re.match(r"(\d+)([A-Za-z]*)", del_)
        ud.append((int(m.group(1)), m.group(2)) if m else (0, del_))
    return ud


def _maal_bilagsnummer(maal_doc):
    """(type, nr) for måldokumentet hvis det er et bilag (fra filnavnet) -
    None hvis det ikke matcher bilagsmønsteret (fx hvis det er kontrakten)."""
    m = kt._bilag_match(maal_doc.path.stem)
    return (kt._btype(m.group(1)), kt.norm_nr(m.group(2))) if m else None


def find_henvisninger_til_maal(kilde_docs, maal_doc, kun_praefikseret=True):
    """Finder alle afsnitsnumre som kilde_docs citerer, OG som eksplicit
    peger på maal_doc - via "bilag X, afsnit Y" hvor X matcher maal_doc's
    eget bilagsnummer, eller "Kontraktens punkt Y" hvis maal_doc er
    kontrakten. Returnerer {normaliseret_nr: [(kildedok, sætning), ...]}.

    kun_praefikseret=False medtager også bare "afsnit Y"/"punkt Y" UDEN
    bilags-/kontraktpræfiks som en henvisning til maal_doc - kun fornuftigt
    når kildedokumentet i det undersøgte område udelukkende/hovedsageligt
    handler om maal_doc's afsnit (ellers blandes interne selvhenvisninger
    ind). Default er den sikre, konservative tilstand.
    """
    maal_bilag = _maal_bilagsnummer(maal_doc)
    maal_er_kontrakt = maal_bilag is None and bool(
        re.search(r"\bkontrakt", maal_doc.path.name, re.IGNORECASE))

    resultat = {}

    def tilfoej(nr, kildedok, ctx):
        resultat.setdefault(kt.norm_nr(nr), []).append((kildedok, ctx))

    for d in kilde_docs:
        navn = d.path.name
        for _, _, tekst, _ in d.paras:
            if not tekst:
                continue
            optaget = []

            def ledig(m):
                return not any(a < m.end() and m.start() < b for a, b in optaget)

            # "bilag X, afsnit Y[, Z og W]" - eksplicit maal-praefikseret
            for m in kt.RE_BILAG_SEKTION.finditer(tekst):
                if kt.er_definition(tekst, m) or not kt.er_bilagsnummer(m.group(2)):
                    continue
                optaget.append((m.start(), m.end()))
                btype, bnr = m.group(1), kt.norm_nr(m.group(2))
                if maal_bilag and kt.bilagstype_match(btype, maal_bilag[0]) and bnr == maal_bilag[1]:
                    ctx = kt.saetning(tekst, m.start(), m.end())
                    tilfoej(m.group(4), navn, ctx)
                    for ekstra in kt.NUM_RE.findall(m.group("hale") or ""):
                        tilfoej(ekstra, navn, ctx)

            # "Kontraktens punkt Y[, Z og W]" - kun relevant når maal er kontrakten,
            # eller (kun_praefikseret=False) bare "afsnit/punkt Y" uden præfiks
            for m in kt.RE_SEKTION.finditer(tekst):
                if not ledig(m) or kt.er_definition(tekst, m):
                    continue
                praefiks = (m.group("praefiks") or "").lower()
                nr1 = m.group("nr1_a") or m.group("nr1_b")
                if not nr1:
                    continue
                nr2 = m.group("nr2_a") or m.group("nr2_b")
                hale = m.group("hale_a") or m.group("hale_b") or ""
                peger_paa_maal = (
                    (maal_er_kontrakt and praefiks.startswith("kontrakt"))
                    or (not kun_praefikseret and not praefiks)
                )
                if not peger_paa_maal:
                    continue
                ctx = kt.saetning(tekst, m.start(), m.end())
                for nr in filter(None, (nr1, nr2, *kt.NUM_RE.findall(hale))):
                    tilfoej(nr, navn, ctx)
    return resultat


def daekningstjek(kilde_docs, maal_doc, fra=None, til=None, kun_praefikseret=True):
    """
    Returnerer (daekket, ikke_daekket, henvisninger) - de to første er
    numerisk sorterede lister af normaliserede afsnitsnumre.

    fra/til: begrænser måldområdet til afsnit hvis øverste kapitelniveau
        ligger i [fra, til] (begge inklusive), fx fra=4, til=6 for "kun
        kapitel 4, 5 og 6". Uden fra/til tjekkes ALLE afsnit i maal_doc -
        kan give meget støj hvis maal_doc har afsnit der aldrig var ment
        at skulle citeres systematisk (fx en indledning).
    """
    maal_sektioner, _ = kt.kortlaeg([maal_doc])
    alle_maal_numre = maal_sektioner[maal_doc.path.name]

    if fra is not None or til is not None:
        def i_omraade(nr):
            m = re.match(r"\d+", nr)
            if not m:
                return False
            top = int(m.group(0))
            return (fra is None or top >= fra) and (til is None or top <= til)
        maal_numre = {nr for nr in alle_maal_numre if i_omraade(nr)}
    else:
        maal_numre = set(alle_maal_numre)

    henvisninger = find_henvisninger_til_maal(kilde_docs, maal_doc, kun_praefikseret)
    daekket = sorted((nr for nr in maal_numre if nr in henvisninger), key=_sort_key)
    ikke_daekket = sorted((nr for nr in maal_numre if nr not in henvisninger), key=_sort_key)
    return daekket, ikke_daekket, henvisninger


# ---------------------------------------------------------------------------
# Rapport
# ---------------------------------------------------------------------------
def skriv_rapport(daekket, ikke_daekket, henvisninger, maal_navn, outfile, intro=""):
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Dækningstjek")
    rr = 1
    if intro:
        ws.cell(row=1, column=1, value=intro).font = Font(italic=True, color="666666")
        rr = 2
    hoved = ["Status", f"Afsnit i {maal_navn}", "Citeret hvor (kildedokument - sætning)"]
    bredder = [18, 16, 110]
    for c, (h, b) in enumerate(zip(hoved, bredder), 1):
        cell = ws.cell(row=rr, column=c, value=h)
        cell.font = Font(bold=True)
        ws.column_dimensions[cell.column_letter].width = b
    rr += 1
    ws.freeze_panes = f"A{rr}"

    for nr in ikke_daekket:
        ws.cell(row=rr, column=1, value="❌ IKKE DÆKKET").fill = FILL_MANGLER
        ws.cell(row=rr, column=2, value=nr)
        ws.cell(row=rr, column=3, value="(nævnes ikke nogen steder i kildedokumenterne)")
        for c in range(1, 4):
            ws.cell(row=rr, column=c).border = THIN
            ws.cell(row=rr, column=c).alignment = Alignment(wrap_text=True, vertical="top")
        rr += 1
    for nr in daekket:
        cites = henvisninger.get(nr, [])
        vis = "; ".join(f"{dok}: {ctx[:120]}" for dok, ctx in cites[:3])
        if len(cites) > 3:
            vis += f" ... (+{len(cites) - 3} mere)"
        ws.cell(row=rr, column=1, value="✅ Dækket").fill = FILL_DAEKKET
        ws.cell(row=rr, column=2, value=nr)
        ws.cell(row=rr, column=3, value=vis)
        for c in range(1, 4):
            ws.cell(row=rr, column=c).border = THIN
            ws.cell(row=rr, column=c).alignment = Alignment(wrap_text=True, vertical="top")
        rr += 1
    if not daekket and not ikke_daekket:
        ws.cell(row=rr, column=2, value="(ingen afsnit fundet i det angivne område)")
    wb.save(outfile)
    return len(daekket), len(ikke_daekket)


# ---------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(
        description="Dækningstjek - opdager afsnit i et måldokument der aldrig citeres fra kildedokumenter")
    ap.add_argument("kilder", nargs="+", metavar="KILDE.docx",
                    help="Dokument(er) der scannes for henvisninger (fx udbudsbetingelserne)")
    ap.add_argument("--maal", required=True, metavar="MAAL.docx",
                    help="Dokumentet hvis afsnit skal tjekkes for dækning (fx kravspecifikationen)")
    ap.add_argument("--fra", type=int, default=None,
                    help="Laveste kapitelnummer der skal tjekkes (fx 4)")
    ap.add_argument("--til", type=int, default=None,
                    help="Højeste kapitelnummer der skal tjekkes (fx 6)")
    ap.add_argument("--tillad-upraefikseret", action="store_true",
                    help="Medtag også bare 'afsnit X'/'punkt X' uden bilags-/kontraktpræfiks "
                         "som en henvisning til måldokumentet - kun fornuftigt hvis "
                         "kildedokumentet i det undersøgte område udelukkende handler om "
                         "måldokumentets afsnit (ellers blandes interne selvhenvisninger ind)")
    ap.add_argument("-o", "--output", default="daekningstjek.xlsx")
    args = ap.parse_args()

    kilde_docs = [ex.Docx(f) for f in args.kilder]
    maal_doc = ex.Docx(args.maal)

    daekket, ikke_daekket, henvisninger = daekningstjek(
        kilde_docs, maal_doc, fra=args.fra, til=args.til,
        kun_praefikseret=not args.tillad_upraefikseret)

    omraade = f"kapitel {args.fra}-{args.til}" if (args.fra or args.til) else "alle afsnit"
    intro = (f"Måldokument: {maal_doc.path.name} | Kildedokument(er): "
            f"{', '.join(d.path.name for d in kilde_docs)} | Område: {omraade}")
    skriv_rapport(daekket, ikke_daekket, henvisninger, maal_doc.path.name, args.output, intro)

    print(f"Dækningstjek gennemført -> {args.output}")
    print(f"  Dækket: {len(daekket)}   Ikke dækket: {len(ikke_daekket)}")
    if ikke_daekket:
        print("  Mangler:", ", ".join(ikke_daekket))


if __name__ == "__main__":
    sys.exit(main())
