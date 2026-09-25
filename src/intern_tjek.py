#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Intern tjek - validering af INTERNE afsnits-/punkthenvisninger i ÉT dokument
=============================================================================
Modsat krydstjek.py (som validerer henvisninger på tværs af en hel
dokumentsamling) tjekker dette modul kun ÉT dokument for sig selv: findes
det nummer en henvisning peger på som et afsnit i samme dokument - og,
hvis henvisningen også nævner et emne ("jf. punkt 18.2 om databehandling",
"punkt 14.1 (Betaling)"), matcher emnet afsnittets faktiske overskrift?

Emne-sammenligningen er simpel ordsammenligning (små bogstaver, hovedord
skal optræde i overskriften) - ingen AI er involveret.

Genbruger nummerkortlægning, regex og hjælpefunktioner fra krydstjek.py
uden at ændre den fil.

Kan køres alene:
    python intern_tjek.py kontrakt.docx -o intern_tjek.xlsx
eller via knappen "Tjek interne henvisninger" i extractor_gui.py.

Kræver: extractor.py og krydstjek.py i samme mappe + openpyxl.
"""

import argparse
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

import extractor as ex
import krydstjek as kt

THIN = Border(*[Side(style="thin")] * 4)
FILLS = {"ugyldig": PatternFill("solid", start_color="F8CBAD"),
         "usikker": PatternFill("solid", start_color="FFE699"),
         "gyldig": PatternFill("solid", start_color="C6E0B4")}
SYMBOL = {"ugyldig": "❌ UGYLDIG", "usikker": "⚠️ USIKKER", "gyldig": "✅ GYLDIG"}

# -- emne-sammenligning -------------------------------------------------
STOPORD = {"og", "i", "for", "den", "det", "en", "et", "som", "der", "til",
           "af", "om", "vedrørende", "angående", "med", "på", "samt",
           "eller", "ikke", "skal", "kan", "ved", "fra", "er", "de", "alle"}

# Emne lige efter en sektionshenvisning: "(...)" eller bindeord som
# "om"/"vedrørende"/"angående" fulgt af et kort emne. Ordfangningen stopper
# ved næste tegnsætning ELLER ved et hjælpeord/bindeord ("skal", "og", ...),
# så kun selve emnet fanges - ikke resten af sætningen ("...om betaling SKAL
# ske rettidigt" -> emnet er "betaling", ikke "betaling skal ske rettidigt").
_EMNE_STOP = (r"(?:skal|kan|vil|bør|må|og|eller|samt|men|fordi|idet|der|som"
             r"|inden|senest|uden|jf\.?)")
RE_EKSTERN_EFTER = re.compile(
    r"^\s*(?:\([^)]*\)\s*)?i\s+"
    r"(?:udbudsbetingelser(?:ne)?|udbudsbetingelse|"
    r"kravspecifikation(?:en)?|kontrakten|bilag\b)",
    re.IGNORECASE)

RE_EMNE = re.compile(
    r"^\s*(?:\((?P<paren>[^)]+)\)|"
    rf"(?:om|vedrørende|angående)\s+(?P<ord>(?:(?!{_EMNE_STOP}\b)[^\s.,;:\n]+\s*){{1,6}}))",
    re.IGNORECASE)


def hovedord(tekst):
    """Betydningsbærende ord i lowercase - korte ord og stopord filtreret."""
    raa = re.findall(r"[a-zæøåA-ZÆØÅ0-9]+", tekst.lower())
    return [o for o in raa if o not in STOPORD and len(o) > 2]


def emne_matcher(emne, overskrift):
    """Sand hvis alle hovedord i emnet optræder i overskriften."""
    ho = hovedord(emne)
    if not ho:
        return True
    overskrift_lower = overskrift.lower()
    return all(o in overskrift_lower for o in ho)


def find_emne(tekst, slut):
    """Kigger lige efter en sektionshenvisnings slut for et nævnt emne."""
    m = RE_EMNE.match(tekst[slut:])
    if not m:
        return None
    emne = (m.group("paren") or m.group("ord") or "").strip()
    return emne or None


def byg_nummer_til_overskrift(docx):
    """{normaliseret nummer: overskriftstekst} for afsnit i dokumentet -
    dækker både Words auto-nummererede overskrifter og manuelt indtastede
    numre (fx "18.1.2 Adgangskontrol").

    To gennemløb med vilje: ægte Word-nummererede overskrifter (label) går
    forud for manuelt matchede numre (RE_MANUELT_NR), da en indholds-
    fortegnelse ofte indeholder tab-separerede linjer som "30.3<tab>
    Servicemål<tab>23" der ELLERS fejlagtigt ville matche RE_MANUELT_NR og
    - fordi TOC'en typisk står først i dokumentet - vinde over den rigtige,
    rene overskriftstekst senere i dokumentet."""
    out = {}
    rest_kandidater = []
    for _, label, tekst, _ in docx.paras:
        if label and label[:1].isdigit():
            out.setdefault(kt.norm_nr(label), tekst)
            continue
        m = kt.RE_MANUELT_NR.match(tekst)
        if m and len(tekst) <= 150:
            rest = tekst[m.end(1):].strip(" .)-\t")
            rest_kandidater.append((kt.norm_nr(m.group(1)), rest))
    for nr, rest in rest_kandidater:
        out.setdefault(nr, rest)
    return out


# -- hovedvalidering ------------------------------------------------------
def validér_internt(docx):
    """Validerer interne afsnits-/punkthenvisninger i ÉT dokument.
    Returnerer (fund, sektioner). fund-tupler: (status, henvisning,
    kontekst, forklaring)."""
    sektioner, _ = kt.kortlaeg([docx])
    numre = sektioner[docx.path.name]
    overskrifter = byg_nummer_til_overskrift(docx)

    fund = []
    for _, _, tekst, _ in docx.paras:
        if not tekst:
            continue
        optaget = []

        def ledig(m):
            return not any(a < m.end() and m.start() < b for a, b in optaget)

        # "Bilag X, punkt Y" henviser til et ANDET dokuments afsnit -
        # marker som optaget så det ikke fejlagtigt tjekkes som internt.
        for m in kt.RE_BILAG_SEKTION.finditer(tekst):
            optaget.append((m.start(), m.end()))

        # Henvisninger til EKSTERN lovgivning ("Udbudslovens § 134 a", "GDPR
        # art. 6") - §-nummeret peger ikke på et afsnit i DETTE dokument, så
        # det ville være forkert at tjekke det som en intern henvisning.
        # Rapporteres som usikker (bør verificeres manuelt), ikke stille
        # udeladt - loven indgår jo aldrig i det vi har at validere imod.
        for m in kt.RE_EKSTERN_LOV.finditer(tekst):
            if not ledig(m) or kt.er_definition(tekst, m):
                continue
            optaget.append((m.start(), m.end()))
            lovnavn = m.group(1) or m.group(4)
            fund.append(("usikker", m.group(0),
                         kt.saetning(tekst, m.start(), m.end()),
                         f"henvisning til ekstern lovgivning ({lovnavn}) - "
                         f"kan ikke tjekkes mod dette dokuments egen "
                         f"afsnitsnummerering; bør verificeres manuelt"))

        for m in kt.RE_SEKTION.finditer(tekst):
            if not ledig(m) or kt.er_definition(tekst, m):
                continue
            if RE_EKSTERN_EFTER.match(tekst[m.end():]):
                optaget.append((m.start(), m.end()))
                continue
            nr1 = m.group("nr1_a") or m.group("nr1_b")
            nr2 = m.group("nr2_a") or m.group("nr2_b")
            hale = m.group("hale_a") or m.group("hale_b") or ""
            ekstra = kt.NUM_RE.findall(hale)
            if not nr1:
                continue
            ctx = kt.saetning(tekst, m.start(), m.end())
            # Emne-sammenligning køres nu for HVERT tal i en opremsning/interval
            # (fx "punkt 18.1, 18.2 og 18.3 (Databehandling)"), ikke kun ved en
            # enkeltstående henvisning. Tidligere blev emne-tjek droppet helt så
            # snart der var mere end ét tal ("giver kun mening for en enkelt
            # henvisning") - men netop den slags opremsninger er hvor en reel
            # fejl (et tal i midten af listen der er forskudt/forkert ift. det
            # emne, listen faktisk beskriver) ellers aldrig bliver opdaget,
            # fordi hvert enkelt tal jo findes og derfor ville blive "gyldig".
            # Risikoen er flere "usikker"-markeringer for lister med en bred
            # fælles overskrift der ikke matcher alle led lige præcist - det er
            # en bevidst, rimelig afvejning: "usikker" beder blot om et
            # menneskeligt kig, det er ikke en hård fejl.
            emne = find_emne(tekst, m.end())

            for nr in filter(None, (nr1, nr2, *ekstra)):
                nr_norm = kt.norm_nr(nr)
                if nr_norm not in numre:
                    fund.append(("ugyldig", m.group(0), ctx,
                                f"afsnit {nr_norm} findes ikke i dokumentet"))
                    continue
                if emne:
                    overskrift = overskrifter.get(nr_norm, "")
                    if overskrift and not emne_matcher(emne, overskrift):
                        fund.append((
                            "usikker", m.group(0), ctx,
                            f"punkt {nr_norm} findes men hedder "
                            f"'{overskrift}', henvisningen kalder det '{emne}'"))
                        continue
                fund.append(("gyldig", m.group(0), ctx,
                            f"afsnit {nr_norm} findes i dokumentet"))
    return fund, numre


# -- rapport ---------------------------------------------------------------
def _skriv_ark(wb, titel, raekker, intro=None):
    ws = wb.create_sheet(titel)
    rr = 1
    if intro:
        ws.cell(row=1, column=1, value=intro).font = Font(italic=True, color="666666")
        rr = 2
    hoved = ["Status", "Henvisning", "Sætning (kontekst)", "Vurdering"]
    bredder = [16, 22, 75, 60]
    for c, (h, b) in enumerate(zip(hoved, bredder), 1):
        cell = ws.cell(row=rr, column=c, value=h)
        cell.font = Font(bold=True)
        ws.column_dimensions[cell.column_letter].width = b
    orden = {"ugyldig": 0, "usikker": 1, "gyldig": 2}
    for status, ref, ktx, forkl in sorted(raekker, key=lambda f: orden[f[0]]):
        rr += 1
        ws.cell(row=rr, column=1, value=SYMBOL[status]).fill = FILLS[status]
        for c, v in enumerate((ref, ktx, forkl), 2):
            ws.cell(row=rr, column=c, value=v)
        for c in range(1, 5):
            ws.cell(row=rr, column=c).border = THIN
            ws.cell(row=rr, column=c).alignment = Alignment(wrap_text=True, vertical="top")
    if rr == (2 if intro else 1):
        ws.cell(row=rr + 1, column=2, value="(ingen)")


def skriv_rapport(fund, outfile):
    wb = Workbook()
    wb.remove(wb.active)
    n = {s: sum(1 for f in fund if f[0] == s) for s in ("gyldig", "ugyldig", "usikker")}
    _skriv_ark(wb, "Til gennemgang",
               [f for f in fund if f[0] in ("ugyldig", "usikker")],
               intro="Henvisninger der kræver menneskelig vurdering - ugyldige øverst.")
    _skriv_ark(wb, "Gyldige", [f for f in fund if f[0] == "gyldig"])
    wb.save(outfile)
    return n


# ---------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(
        description="Validér interne afsnits-/punkthenvisninger i ét dokument")
    ap.add_argument("fil", help="Word-dokument (.docx)")
    ap.add_argument("-o", "--output", default="intern_tjek.xlsx")
    args = ap.parse_args()

    docx = ex.Docx(args.fil)
    fund, numre = validér_internt(docx)
    n = skriv_rapport(fund, args.output)
    print(f"Intern tjek gennemfort -> {args.output}")
    print(f"  Ugyldige: {n['ugyldig']}   Usikre: {n['usikker']}   Gyldige: {n['gyldig']}")


if __name__ == "__main__":
    sys.exit(main())
