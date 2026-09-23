"""Tegner blueprintet i Rambølls blueprint-template ved at "stemple" templatens
egne figurer: paletten over diaset (pilefigurer i farvekategorier), en
milepælsgruppe (linje + Mx-boks + navneboks) og symbolerne i signaturforklaringen.

Figurerne findes via deres shape-id i templaten (se TEMPLATE_KONFIG). Anbefaling:
navngiv stempel-figurerne i PowerPoints markeringsrude, fx "STEMPEL_fase1",
og slå op på navn i stedet for id, så templaten kan redigeres frit.
"""
from __future__ import annotations

import copy
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Inches

from .model import BANER, Blueprint

TEMPLATE_KONFIG = {
    "dias": 1,                       # 0-baseret: dias 2 = symboler i banerne
    "milepael_gruppe": 16411,        # Group 16410: linje + "Delleverance-prøve" + "Mx"
    "aktivitet": {                   # farvekategori -> pentagon i paletten
        "fase1": 78, "fase1_2l": 83, "fase2": 89, "fase2_2l": 90,
        "drift": 140, "drift_2l": 86, "videreudvikling": 87, "videreudvikling_2l": 88,
        "governance": 91, "ophoer": 107, "uddannelse": 98,
    },
    "symbol": {                      # retsvirkning -> figur i signaturforklaringen
        "betaling": 16397, "ophævelse": 16398, "bod": 16399, "ibrugtagning": 16400,
        "udtrædelsesadgang": 117, "overtagelse": 121, "bonus": 111,
    },
    "x_start": 1.72, "x_slut": 12.14,  # tommer: kontraktindgåelse .. kontraktophør
    "baner_top": {"Etableringsprojekt": 0.46, "Løbende ydelser": 3.53, "Optioner": 5.14,
                  "Governance": 5.46, "Kundens ydelser": 5.78},
    "symbol_top": 2.2,               # første symbol under aktiviteterne i Etableringsprojekt
    "raekkehoejde": 0.36,
}

def _figur(slide, shape_id):
    for sh in slide.shapes:
        if sh.shape_id == shape_id:
            return sh
    raise KeyError(f"Figur med id {shape_id} findes ikke i templaten")


def _klon(slide, kilde, left=None, top=None, width=None, tekst=None):
    el = copy.deepcopy(kilde._element)
    slide.shapes._spTree.insert_element_before(el, "p:extLst")
    ny = slide.shapes[-1]
    if left is not None:
        ny.left = Emu(int(left))
    if top is not None:
        ny.top = Emu(int(top))
    if width is not None:
        ny.width = Emu(int(width))
    if tekst is not None and ny.has_text_frame:
        _saet_tekst(ny, tekst)
    return ny


def _saet_tekst(shape, tekst: str):
    """Bevarer formateringen fra første run og skriver ny tekst."""
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    runs = p0.runs
    if runs:
        runs[0].text = tekst
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        p0.text = tekst


def _ryd_laerred(slide):
    """Fjerner eksempelindholdet på lærredet; palette og signaturforklaring bevares."""
    y_min, y_max = Inches(0.1), Inches(6.8)
    slet = []
    for sh in slide.shapes:
        tekst = sh.text_frame.text if sh.has_text_frame else ""
        er_milepael = sh.shape_type == 6 and any(
            c.has_text_frame and c.text_frame.text in ("Mx", "Kontrakt-indgåelse")
            for c in sh.shapes)
        er_indhold = sh.shape_type == 1 and y_min < sh.top < y_max and sh.left > Inches(1.1) \
            and tekst not in ("Milepæle",)
        if er_milepael or er_indhold:
            slet.append(sh)
    for sh in slet:
        sh._element.getparent().remove(sh._element)


def _behold_kun_dias(prs, indeks: int):
    sld_ids = prs.slides._sldIdLst
    for i, sld in reversed(list(enumerate(list(sld_ids)))):
        if i != indeks:
            prs.part.drop_rel(sld.rId)
            sld_ids.remove(sld)


def _titel(kunde: str, udbud: str) -> str:
    dele = [d for d in (udbud, kunde) if d]
    return "Kontraktens blueprint – " + " – ".join(dele) if dele else "Kontraktens blueprint"


def generer(bp: Blueprint, template: str | Path, ud: str | Path, konfig: dict = TEMPLATE_KONFIG,
            vis_typer: tuple[str, ...] = ("kontraktindgåelse", "faseafslutning", "milepæl", "prøve", "ophør"),
            kunde: str = "", udbud: str = ""):
    """Tegner blueprintet. Frister vises kun, hvis "frist" er med i vis_typer.
    Angives kunde og/eller udbud, saettes de i dias-titlen; ellers bruges templatens egen titel."""
    prs = Presentation(str(template))
    slide = prs.slides[konfig["dias"]]
    if (kunde or udbud) and slide.shapes.title is not None:
        _saet_tekst(slide.shapes.title, _titel(kunde, udbud))
    stempler = {
        "milepael": _figur(slide, konfig["milepael_gruppe"]),
        "aktivitet": {k: _figur(slide, v) for k, v in konfig["aktivitet"].items()},
        "symbol": {k: _figur(slide, v) for k, v in konfig["symbol"].items()},
    }
    # Klon stemplerne fra lærredet før det ryddes
    mp_stempel = copy.deepcopy(stempler["milepael"]._element)
    _ryd_laerred(slide)
    slide.shapes._spTree.insert_element_before(mp_stempel, "p:extLst")
    mp_kilde = slide.shapes[-1]

    ms = sorted((m for m in bp.milepaele if m.type in vis_typer), key=lambda m: m.raekkefoelge)
    x0, x1 = Inches(konfig["x_start"]), Inches(konfig["x_slut"])
    skridt = (x1 - x0) / max(len(ms) - 1, 1)
    x_for = {m.id: x0 + i * skridt for i, m in enumerate(ms)}

    # Milepæle: gruppens linje ligger 0,59" inde i gruppen
    linje_offset = Inches(0.59)
    for m in ms:
        g = _klon(slide, mp_kilde, left=x_for[m.id] - linje_offset)
        for c in g.shapes:
            if c.has_text_frame and c.text_frame.text == "Mx":
                _saet_tekst(c, m.id)
            elif c.has_text_frame and c.text_frame.text:
                _saet_tekst(c, m.etiket or m.navn)
    mp_kilde._element.getparent().remove(mp_kilde._element)

    # Aktiviteter
    for a in bp.aktiviteter:
        if a.fra not in x_for or a.til not in x_for:
            continue
        kilde = stempler["aktivitet"].get(a.farvekategori) or stempler["aktivitet"]["fase1"]
        venstre, hoejre = x_for[a.fra], x_for[a.til]
        top = Inches(konfig["baner_top"][a.bane] + a.raekke * konfig["raekkehoejde"])
        _klon(slide, kilde, left=venstre, top=top, width=max(hoejre - venstre, Inches(0.5)),
              tekst=a.tekst)

    # Retsvirkninger: symboler stables under hinanden ved milepælens linje
    stak: dict[str, int] = {}
    tegnet: set[tuple[str, str]] = set()
    for r in bp.retsvirkninger:
        kilde = stempler["symbol"].get(r.type)
        if kilde is None or r.milepael not in x_for or (r.milepael, r.type) in tegnet:
            continue
        tegnet.add((r.milepael, r.type))
        n = stak.get(r.milepael, 0)
        stak[r.milepael] = n + 1
        kol, rk = divmod(n, 4)          # maks. 4 symboler pr. kolonne
        _klon(slide, kilde, left=x_for[r.milepael] - kilde.width // 2 + Inches(0.28) * kol,
              top=Inches(konfig["symbol_top"] + rk * 0.25))

    _behold_kun_dias(prs, konfig["dias"])
    prs.save(str(ud))
    return ud
