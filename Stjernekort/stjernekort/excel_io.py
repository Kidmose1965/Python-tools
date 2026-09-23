"""Excel-arket er master: skriv et Blueprint til .xlsx og læs det tilbage."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .model import BANER, RETSVIRKNINGER, Aktivitet, Blueprint, Kilde, Milepael, Retsvirkning

FANER = {
    "Milepæle": ["ID", "Navn", "Etiket", "Type", "Rækkefølge", "Option", "Relativ tid", "Dato", "Kilder"],
    "Aktiviteter": ["Bane", "Tekst", "Fra", "Til", "Farvekategori", "Række"],
    "Retsvirkninger": ["Milepæl", "Type", "Detalje", "Hjemmel", "Citat", "Oprindelse"],
    "Henvisninger": ["Fra", "Til dokument", "Til punkt", "Fundet"],
}


def skriv(bp: Blueprint, sti: str | Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    for navn, kol in FANER.items():
        ws = wb.create_sheet(navn)
        ws.append(kol)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E79")
        ws.freeze_panes = "A2"
    ws = wb["Milepæle"]
    for m in bp.milepaele:
        ws.append([m.id, m.navn, m.etiket, m.type, m.raekkefoelge, "ja" if m.option else "nej",
                   m.relativ_tid, m.dato, "; ".join(str(k) for k in m.kilder[:6])])
    ws = wb["Aktiviteter"]
    for a in bp.aktiviteter:
        ws.append([a.bane, a.tekst, a.fra, a.til, a.farvekategori, a.raekke])
    dv = DataValidation(type="list", formula1='"' + ",".join(BANER) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("A2:A500")
    ws = wb["Retsvirkninger"]
    for r in bp.retsvirkninger:
        ws.append([r.milepael, r.type, r.detalje, str(r.kilde) if r.kilde else "",
                   r.kilde.citat if r.kilde else "", r.oprindelse])
    dv2 = DataValidation(type="list", formula1='"' + ",".join(RETSVIRKNINGER) + '"', allow_blank=True)
    ws.add_data_validation(dv2)
    dv2.add("B2:B500")
    ws = wb["Henvisninger"]
    for h in bp.henvisninger:
        ws.append([str(h.fra), h.til_dokument, h.til_punkt,
                   {True: "ja", False: "NEJ", None: "dokument mangler"}[h.fundet]])
    for ws in wb.worksheets:
        for kol in ws.columns:
            ws.column_dimensions[kol[0].column_letter].width = min(
                60, max(10, *(len(str(c.value or "")) for c in kol[:50])) + 2)
    sti = Path(sti)
    wb.save(sti)
    return sti


def _kilde(tekst: str) -> Kilde | None:
    if not tekst:
        return None
    dok, _, punkt = tekst.partition(" pkt. ")
    return Kilde(dok, punkt)


def laes(sti: str | Path) -> Blueprint:
    wb = load_workbook(sti, data_only=True)
    bp = Blueprint(navn=Path(sti).stem)
    for r in wb["Milepæle"].iter_rows(min_row=2, values_only=True):
        if r[0]:
            bp.milepaele.append(Milepael(
                id=str(r[0]), navn=r[1] or "", etiket=r[2] or "", type=r[3] or "",
                raekkefoelge=int(r[4] or 0), option=(r[5] or "").lower() == "ja",
                relativ_tid=r[6] or "", dato=str(r[7] or "")))
    for r in wb["Aktiviteter"].iter_rows(min_row=2, values_only=True):
        if r[0]:
            bp.aktiviteter.append(Aktivitet(r[0], r[1] or "", str(r[2]), str(r[3]),
                                            r[4] or "fase1", int(r[5] or 0)))
    for r in wb["Retsvirkninger"].iter_rows(min_row=2, values_only=True):
        if r[0]:
            bp.retsvirkninger.append(Retsvirkning(str(r[0]), r[1] or "", r[2] or "",
                                                  _kilde(r[3] or ""), r[5] or ""))
    return bp
